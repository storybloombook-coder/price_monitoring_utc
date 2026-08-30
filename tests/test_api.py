from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from price_monitor_v4.app import apply_marketplace_resolution, create_app, normalize_marketplace_task, normalize_shop_result
from price_monitor_v4.browser_bridge import EXTENSION_ORIGIN
from price_monitor_v4.catalog import CatalogStore
from price_monitor_v4.config import Settings


def test_success_requires_a_priced_marketplace_offer_with_a_link() -> None:
    empty = normalize_marketplace_task({"marketplace": "Kaina24", "status": "SUCCESS"})
    assert empty["status"] == "NOT_FOUND"

    priced = normalize_marketplace_task({
        "marketplace": "Hinnavaatlus",
        "status": "SUCCESS",
        "cheapest_in_stock": {"price_eur": 199, "url": "https://shop.example/?a=1&amp;b=2"},
    })
    assert priced["status"] == "SUCCESS"
    assert priced["cheapest_in_stock"]["url"] == "https://shop.example/?a=1&b=2"

    legacy_preorder = normalize_marketplace_task({
        "marketplace": "Kaina24", "status": "SUCCESS",
        "cheapest_preorder": {"price_eur": 180, "url": "https://shop.example/preorder"},
    })
    assert legacy_preorder["status"] == "SUCCESS"
    assert legacy_preorder["cheapest_pre_order"]["price_eur"] == 180

    corrected = apply_marketplace_resolution(priced, {
        "status": "NOT_FOUND", "price_eur": None, "availability": None,
        "seller_name": None, "product_url": None, "decided_at": "2026-08-29T00:00:00+00:00",
    })
    assert corrected["status"] == "NOT_FOUND"
    assert corrected["cheapest_in_stock"] is None


def test_old_search_page_price_is_discarded() -> None:
    result = normalize_shop_result({
        "status": "SUCCESS",
        "price_eur": 99,
        "collection_method": "direct",
        "product_url": "https://www.euronics.ee/otsing/25g64",
        "search_url": "https://www.euronics.ee/otsing/25G64",
    })
    assert result["status"] == "NOT_FOUND"
    assert result["price_eur"] is None


def test_catalog_api_and_v4_health(tmp_path: Path) -> None:
    settings = Settings.load(tmp_path)
    settings = Settings(
        **{
            **settings.__dict__,
            "legacy_enabled": False,
            "open_browser": False,
        }
    )
    store = CatalogStore(settings.catalog_database)
    app = create_app(settings=settings, store=store)

    with TestClient(app) as client:
        health = client.get("/health")
        assert health.status_code == 200
        assert health.json()["version"] == "4.2.4"
        assert health.json()["browser_bridge"]["connected"] is False
        assert health.json()["polite_monitoring"]["per_shop_concurrency"] == 1
        assert health.json()["polite_monitoring"]["browser_concurrency"] == 2
        assert health.json()["polite_monitoring"]["negative_cache_ttl_seconds"] == 21_600
        assert client.post("/browser-bridge/heartbeat").status_code == 403
        heartbeat = client.post("/browser-bridge/heartbeat", headers={"origin": EXTENSION_ORIGIN})
        assert heartbeat.status_code == 200
        assert heartbeat.json()["connected"] is True
        with client.websocket_connect("/browser-bridge/ws", headers={"origin": EXTENSION_ORIGIN}) as socket:
            assert socket.receive_json()["type"] == "ready"
            assert client.get("/browser-bridge/status").json()["transport"] == "websocket"
            socket.send_json({"type": "heartbeat"})
            assert socket.receive_json()["type"] == "ack"

        created = client.post(
            "/catalog/items/source",
            json={"model": "55P7L", "source_sheets": ["TV"]},
        )
        assert created.status_code == 201
        item_id = created.json()["id"]

        paused = client.patch(f"/catalog/items/{item_id}", json={"paused": True})
        assert paused.json()["state"] == "paused"

        trashed = client.delete(f"/catalog/items/{item_id}")
        assert trashed.json()["state"] == "trash"
        assert client.get("/catalog/items?kind=source&scope=active").json() == []

        restored = client.post(f"/catalog/items/{item_id}/restore")
        assert restored.json()["state"] == "paused"

        stock_item = client.post(
            "/catalog/items/stock", json={"nomenclature": "TCL TV", "model": "55T7B"}
        ).json()
        promoted = client.post(f"/catalog/items/{stock_item['id']}/monitor")
        assert promoted.status_code == 200
        assert promoted.json()["model"] == "55T7B"

        client.delete(f"/catalog/items/{item_id}")
        permanent = client.delete(f"/catalog/items/{item_id}/permanent")
        assert permanent.status_code == 204

        sources = client.get("/sources").json()
        assert {item["name"] for item in sources if item["kind"] == "shop"} == {
            "Senukai", "Bite", "Varle", "Elesen", "Elisa", "Euronics", "RDE", "Smartech"
        }
        method = client.patch("/sources/varle", json={"collection_method": "playwright"})
        assert method.status_code == 200
        assert method.json()["collection_method"] == "playwright"
        assert client.patch("/sources/varle", json={"collection_method": "legacy"}).status_code == 400
        client.patch("/sources/master/marketplace", json={"enabled": False})
        client.patch("/sources/master/shop", json={"enabled": False})
        assert client.post("/runs", json={"mode": "turbo"}).status_code == 400
        started = client.post("/runs", json={"mode": "quick"})
        assert started.status_code == 200
        assert started.json()["mode"] == "quick"
        run = client.get(f"/runs/{started.json()['run_id']}").json()
        assert run["status"] == "COMPLETE"
        assert run["run_mode"] == "quick"
        assert run["execution"]["mode"] == "quick"
        assert run["tasks"] == []
        assert run["shop_results"] == []
        export_path = Path(run["export_path"])
        assert export_path.exists()
        exported = client.get(f"/runs/{started.json()['run_id']}/export")
        assert exported.status_code == 200
        assert exported.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        workbook = load_workbook(export_path, read_only=True)
        headers = [cell.value for cell in next(workbook["Monitoring summary"].iter_rows())]
        assert headers[3:12] == [
            "Lowest pre-order", "Senukai", "Bite", "Varle", "Elesen",
            "Elisa", "Euronics", "RDE", "Smartech",
        ]

        source = next(item for item in store.list_items("source", "active") if item["model"] == "55T7B")
        bite = next(item for item in store.list_sources() if item["key"] == "bite")
        store.register_monitoring_session("manual-stop-test", False, [])
        store.start_shop_run("manual-stop-test", [source], [bite])
        stopped = client.post("/runs/manual-stop-test/stop")
        assert stopped.status_code == 200
        assert stopped.json()["status"] == "INCOMPLETE"
        assert stopped.json()["shop_results"][0]["status"] == "INCOMPLETE"
        resolved = client.post(
            f"/runs/manual-stop-test/shops/{source['id']}/bite/resolve",
            json={"status": "NOT_FOUND"},
        )
        assert resolved.status_code == 200
        assert resolved.json()["status"] == "NOT_FOUND"
        assert resolved.json()["collection_method"] == "manual"

        cleared = client.post("/runs/manual-stop-test/clear")
        assert cleared.status_code == 200
        assert cleared.json()["cleared"] is True
        assert cleared.json()["tasks"] == []
        assert cleared.json()["shop_results"] == []
        assert client.get("/runs/latest").json()["cleared"] is True
        history = client.get("/monitoring-history").json()
        assert history[0]["run_id"] == "manual-stop-test"
        assert history[0]["run_mode"] == "balanced"
        invalid_link = client.post(
            f"/runs/manual-stop-test/shops/{source['id']}/bite/link",
            json={"url": "https://example.com/not-bite"},
        )
        assert invalid_link.status_code == 400
