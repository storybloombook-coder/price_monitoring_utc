import asyncio
import json
from pathlib import Path
import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from price_monitor_v5.catalog import CatalogStore
from price_monitor_v5.offers import parse_page, exact_model, seller_key
from price_monitor_v5.sources import marketplace_url, search_url
from price_monitor_v5.monitor import MarketplaceMonitor
from price_monitor_v5.browser_bridge import BrowserBridge
from price_monitor_v5.config import Settings
from price_monitor_v5.app import create_app
from price_monitor_v5.exporter import write_monitoring_export
from price_monitor_v5 import __version__


PRODUCT_URL = "https://www.hinnavaatlus.ee/6366101/tcl-25-lcd-25g64/"
PRODUCT = '''<h1>Monitor TCL 25" LCD 25G64</h1>
<table><tr class="offer"><td><a class="name" href="/dealer/393-elisa">Elisa Eesti</a></td><td class="in-stock">Laos</td><td><button class="offer-price">199,00 €</button><button>5.53 € kuus</button></td><td><a href="https://www.elisa.ee/private">Poodi</a></td></tr>
<tr class="offer"><td><a class="name">Other Shop</a></td><td class="in-stock">Laos</td><td class="offer-price">249,00 €</td></tr>
<tr class="offer"><td><a class="name">Preorder shop</a></td><td class="in-stock">Ettetellimine</td><td class="offer-price">80,00 €</td></tr></table>'''


@pytest.fixture
def store(tmp_path):
    return CatalogStore(tmp_path / "db.sqlite3")


def test_parser_all_sellers_not_installments():
    parsed = parse_page("hinnavaatlus","25G64",PRODUCT,PRODUCT_URL)
    assert len(parsed["offers"]) == 3
    assert parsed["offers"][0]["price_eur"] == 199
    assert parsed["offers"][0]["url"] == PRODUCT_URL
    assert parsed["offers"][2]["availability"] == "PRE_ORDER"
    assert not parse_page("hinnavaatlus","27G64",PRODUCT,PRODUCT_URL)["offers"]


@pytest.mark.parametrize("url",["https://www.varle.lt/p/25g64", "https://hinnavaatlus.ee.evil.test/p/", "http://www.hinnavaatlus.ee/", "https://www.hinnavaatlus.ee/redirect?url=https://varle.lt/", "https://www.hinnavaatlus.ee:8080/", "https://user:secret@www.hinnavaatlus.ee/", "https://www.hinnavaatlus.ee/out/23"])
def test_url_guard(url):
    with pytest.raises(ValueError):
        marketplace_url("hinnavaatlus",url)


def test_exact_and_aliases():
    assert exact_model("25G64","TCL (25G64)")
    assert not exact_model("25G64","TCL 27G64")
    assert not exact_model("75C6K","TCL 75C6K Pro")
    assert not exact_model("S55HE","TCL S55H")
    assert seller_key("Bitė.lt") == "bite"
    assert seller_key("Elisa Eesti") == "elisa"
    assert seller_key("Euronics.hu") is None
    assert seller_key("Euronics","salidzini") is None
    assert seller_key("Euronics.ee","salidzini") == "euronics"
    assert "TCL+25G64" in search_url("salidzini","25G64")


def test_auto_stock_preserves_pause_duplicates_and_zero(store,tmp_path):
    paused = store.create_item("source",{"model":"25G64","paused":True})
    book=Workbook(); sheet=book.active
    sheet.append([None,"Warehouse A"])
    sheet.append([None,"TCL 25G64",0,100])
    sheet.append([None,"TCL 65P7K",0,300])
    sheet.append([None,"Warehouse B"])
    sheet.append([None,"TCL 65P7K",3,300])
    sheet.append([None,"Cable 123",2,4])
    file=tmp_path/"stock.xlsx"; book.save(file)
    result=store.import_stock_workbook(file,"stock.xlsx")
    assert result["enrolled"]==1 and result["needs_review"]==1
    assert len(store.list_items("source"))==2
    assert store.get_item(paused["id"])["paused"]
    store.begin_run("one","deep")
    assert len(store.checks("one"))==3
    assert all(t["source_model"]=="65P7K" for t in store.checks("one"))
    store.import_stock_workbook(file,"stock.xlsx")
    assert len(store.list_items("source"))==2
    book2=Workbook(); book2.active.append([None,"TCL 55T7B",0,200]); book2.save(file)
    store.import_stock_workbook(file,"new.xlsx")
    assert len(store.list_items("source"))==3


def test_network_only_marketplaces_and_run_extrema(store,tmp_path):
    async def scenario():
        store.create_item("source",{"model":"TCL 25G64"})
        requests=[]
        def handler(request):
            requests.append(str(request.url))
            if request.url.host=="www.hinnavaatlus.ee":
                body=PRODUCT if "/6366101/" in request.url.path else '<a href="/6366101/tcl-25-lcd-25g64/">TCL 25G64</a><a href="https://www.elisa.ee/25G64">TCL 25G64</a>'
                return httpx.Response(200,text=body)
            return httpx.Response(403,text="Captcha")
        monitor=MarketplaceMonitor(store,BrowserBridge(),delay=0,transport=httpx.MockTransport(handler))
        store.begin_run("test","deep"); monitor.launch("test")
        await asyncio.gather(*list(monitor.jobs.values()))
        assert all(any(domain in u for domain in ("kaina24.lt","salidzini.lv","hinnavaatlus.ee")) for u in requests)
        run=store.run("test")
        task=next(t for t in run["tasks"] if t["marketplace_key"]=="hinnavaatlus")
        assert task["cheapest_in_stock"]["price_eur"]==199
        assert task["highest_in_stock"]["price_eur"]==249
        assert task["cheapest_pre_order"]["price_eur"]==80
        assert next(s for s in run["shop_results"] if s["shop_key"]=="elisa")["price_eur"]==199
        assert next(s for s in run["shop_results"] if s["shop_key"]=="rde")["status"]=="UNVERIFIED"
        assert run["status"]=="COMPLETE"
        assert next(t for t in run["tasks"] if t["marketplace_key"]=="salidzini")["status"]=="ACTION_REQUIRED"
        file=tmp_path/"export.xlsx"; write_monitoring_export(file,run)
        workbook=load_workbook(file)
        assert workbook.active["B1"].value=="Marketplace min price"
        assert workbook.active["C1"].value=="Marketplace max price"
        assert workbook.active.page_setup.fitToWidth==1
        assert workbook["All seller offers"].max_row==4
        assert workbook["All seller offers"]["F2"].hyperlink.target==PRODUCT_URL
        assert not monitor.jobs
    asyncio.run(scenario())


def test_redirect_rejected_before_shop_request(store):
    async def scenario():
        requests=[]
        def handler(request):
            requests.append(str(request.url))
            return httpx.Response(302,headers={"location":"https://www.varle.lt/25G64"})
        monitor=MarketplaceMonitor(store,BrowserBridge(),delay=0,transport=httpx.MockTransport(handler))
        async with httpx.AsyncClient(transport=monitor.transport) as client:
            with pytest.raises(ValueError):
                await monitor.fetch(client,"hinnavaatlus",PRODUCT_URL)
        assert len(requests)==1
    asyncio.run(scenario())


def test_manual_offers_correction_not_found_history(store):
    async def scenario():
        item=store.create_item("source",{"model":"25G64"})
        store.begin_run("one","deep")
        monitor=MarketplaceMonitor(store,BrowserBridge())
        base={"status":"SUCCESS","seller_name":"Elisa Eesti","price_eur":199,"availability":"IN_STOCK","product_url":PRODUCT_URL}
        await monitor.resolve("one",item["id"],"hinnavaatlus",base)
        await monitor.resolve("one",item["id"],"hinnavaatlus",{**base,"seller_name":"Other","price_eur":299,"all_offers_reviewed":True})
        task=store.check("one",item["id"],"hinnavaatlus")
        assert len(task["offers"])==2 and task["coverage"]=="complete"
        await monitor.resolve("one",item["id"],"hinnavaatlus",{**base,"offer_index":0,"price_eur":189})
        assert store.check("one",item["id"],"hinnavaatlus")["offers"][1]["price_eur"]==299
        store.begin_run("two","balanced")
        next_task=store.check("two",item["id"],"hinnavaatlus")
        assert not next_task["cached"] and next_task["previous_manual_resolution"]["price_eur"]==189
        await monitor.resolve("two",item["id"],"hinnavaatlus",{"status":"NOT_FOUND"})
        assert not store.check("two",item["id"],"hinnavaatlus")["offers"]
    asyncio.run(scenario())


def test_api_disables_shops_and_cleared_run_stays_empty(store,tmp_path):
    app=create_app(Settings.load(tmp_path),store=store)
    with TestClient(app) as client:
        assert client.get("/health").json()["version"]==__version__
        item=client.post("/catalog/items/stock",json={"nomenclature":"TCL 25G64","quantity":0}).json()
        assert len(client.get("/catalog/items?kind=source").json())==1
        assert client.post("/sources/varle/test",json={}).status_code==404
        assert client.patch("/sources/varle",json={"collection_method":"direct"}).status_code==400
        store.begin_run("one","deep")
        assert client.post("/runs/one/clear").json()["cleared"]
        assert client.get("/runs/one").json()["tasks"]==[]
        assert client.get("/monitoring-history").json()[0]["status"]=="CLEARED"


def test_705_checks_only_three_protected_requests(store):
    async def scenario():
        from price_monitor_v5.catalog import utc_now
        with store.connect() as db:
            db.executemany("INSERT INTO catalog_items(kind,model,canonical_model,source_sheets,created_at,updated_at) VALUES('source',?,?,'[\"TV\"]',?,?)",
                           [(f"55P{i}K",f"55P{i}K",utc_now(),utc_now()) for i in range(235)])
        requests=[]
        def handler(request):
            requests.append(request.url.host)
            return httpx.Response(403)
        monitor=MarketplaceMonitor(store,BrowserBridge(),delay=0,transport=httpx.MockTransport(handler))
        store.begin_run("load","deep")
        assert len(store.checks("load"))==705
        monitor.launch("load")
        await asyncio.gather(*list(monitor.jobs.values()))
        assert len(requests)==3
        run=store.run("load")
        assert run["status"]=="COMPLETE"
        assert all(t["status"]=="ACTION_REQUIRED" for t in run["tasks"])
        assert not monitor.jobs
    asyncio.run(scenario())


def test_hard_stop_no_late_overwrite(store):
    async def scenario():
        store.create_item("source",{"model":"25G64"})
        began=asyncio.Event()
        async def handler(request):
            began.set()
            await asyncio.Event().wait()
        monitor=MarketplaceMonitor(store,BrowserBridge(),delay=0,transport=httpx.MockTransport(handler))
        store.begin_run("stop","deep"); monitor.launch("stop")
        await began.wait()
        await asyncio.wait_for(monitor.stop_run("stop"),timeout=1)
        assert all(t["status"]=="INCOMPLETE" for t in store.checks("stop"))
        assert not monitor.jobs
    asyncio.run(scenario())


def test_cache_ttl_and_manual_correction_invalidation(store):
    from price_monitor_v5.catalog import utc_now
    from datetime import UTC,datetime,timedelta
    async def scenario():
        item=store.create_item("source",{"model":"25G64"})
        store.begin_run("first","deep")
        task=store.check("first",item["id"],"hinnavaatlus")
        task.update(status="SUCCESS",offers=parse_page("hinnavaatlus","25G64",PRODUCT,PRODUCT_URL)["offers"],finished_at=utc_now(),coverage="complete",collection_method="marketplace HTML")
        store.save_check(task)
        assert store.cached_check("25G64","hinnavaatlus",4)
        task["finished_at"]=(datetime.now(UTC)-timedelta(hours=5)).isoformat(); store.save_check(task)
        assert not store.cached_check("25G64","hinnavaatlus",4)
        assert store.cached_check("25G64","hinnavaatlus",12)
        monitor=MarketplaceMonitor(store,BrowserBridge())
        await monitor.resolve("first",item["id"],"hinnavaatlus",{"status":"NOT_FOUND"})
        assert not store.cached_check("25G64","hinnavaatlus",12)
    asyncio.run(scenario())


@pytest.mark.parametrize("key,html,url",[
    ("kaina24",'<h1>TCL 25G64</h1><div class="shop-row"><span class="shop-name">Bitė.lt</span><span class="price">180,00 €</span><span>Turime</span></div>',"https://www.kaina24.lt/p/tcl-25g64/"),
    ("salidzini",'<div class="item_block"><span class="item_name">TCL 25G64</span><span class="item_shop_name">Other.lv</span><span class="item_price">189,99 €</span><span>Ir noliktavā</span></div>',"https://www.salidzini.lv/cena?q=TCL+25G64")])
def test_other_marketplace_rows_are_conservative(key,html,url):
    parsed=parse_page(key,"25G64",html,url)
    assert len(parsed["offers"])==1 and parsed["partial"]
    assert parsed["offers"][0]["availability"]=="IN_STOCK"
    assert parsed["offers"][0]["url"]==url


def test_migration_does_not_copy_prices_or_change_v4(store,tmp_path):
    from price_monitor_v4.catalog import CatalogStore as V4
    old=V4(tmp_path/"v4.sqlite3")
    item=old.create_item("source",{"model":"25G64","paused":True,"shop_links":{"varle":"https://www.varle.lt/p/"}})
    old.create_item("stock",{"model":"25G64","nomenclature":"TCL 25G64","quantity":0})
    store.import_v4_catalog(old.database)
    assert store.get_item(item["id"])["paused"]
    assert not store.get_item(item["id"])["shop_links"]
    assert old.get_item(item["id"])["shop_links"]
    assert not store.latest_monitoring_session()
