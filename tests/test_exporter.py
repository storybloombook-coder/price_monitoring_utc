from pathlib import Path

import pytest
from openpyxl import load_workbook

from price_monitor_v4.exporter import write_monitoring_export


def test_direct_shop_price_participates_in_lowest_in_stock(tmp_path: Path) -> None:
    output = tmp_path / "monitoring.xlsx"
    run = {
        "tasks": [
            {
                "source_model": "25G64",
                "canonical_model": "25G64",
                "marketplace": "Hinnavaatlus",
                "status": "SUCCESS",
                "cheapest_in_stock": {"price_eur": 199, "store": "Elisa"},
                "cheapest_pre_order": None,
                "stock_quantity": 3,
                "stock_unit_cost_eur": 150,
            }
        ],
        "shop_results": [
            {
                "model": "25G64",
                "shop_key": "bite",
                "shop_name": "Bite",
                "status": "SUCCESS",
                "availability": "IN_STOCK",
                "price_eur": 176.4,
                "product_url": "https://www.bite.lt/example",
            }
        ],
    }

    write_monitoring_export(output, run, [{"model": "25G64", "canonical_model": "25G64"}])

    workbook = load_workbook(output, data_only=True, read_only=True)
    row = next(workbook["Monitoring summary"].iter_rows(min_row=2, values_only=True))
    assert row[2] == 176.4
    assert row[14] == pytest.approx(26.4)
