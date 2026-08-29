from pathlib import Path
import sqlite3

import pytest
from openpyxl import Workbook

from price_monitor_v4.catalog import CatalogStore


def make_source_workbook(path: Path) -> None:
    workbook = Workbook()
    tv = workbook.active
    tv.title = "TV"
    tv.append(["Customer", "Model"])
    tv.append(["UTC", "55P7L"])
    sb = workbook.create_sheet("SB")
    sb.append(["Model"])
    sb.append(["Q65H"])
    workbook.create_sheet("Monitors").append(["Model"])
    workbook.save(path)


def make_stock_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Stock"
    sheet.append([None, "Main Warehouse", None, None])
    sheet.append([None, "TV TCL 55P7L", 3, 310.5])
    sheet.append([None, "Unknown accessory", 7, 12])
    workbook.save(path)


def test_import_edit_pause_trash_and_restore(tmp_path: Path) -> None:
    store = CatalogStore(tmp_path / "catalog.sqlite3")
    source_path = tmp_path / "source.xlsx"
    stock_path = tmp_path / "stock.xlsx"
    make_source_workbook(source_path)
    make_stock_workbook(stock_path)

    source_result = store.import_source_workbook(source_path, "source.xlsx")
    stock_result = store.import_stock_workbook(stock_path, "stock.xlsx")

    assert source_result == {"imported": 2, "inserted": 2, "updated": 0}
    assert stock_result == {"imported": 2, "matched": 1}
    source = store.list_items("source")[0]
    edited = store.update_item(source["id"], {"model": "55P7L Pro", "paused": True})
    assert edited["model"] == "55P7L Pro"
    assert edited["paused"] is True
    assert edited["origin"] == "manual"

    trashed = store.trash_item(source["id"])
    assert trashed["state"] == "trash"
    assert source["id"] not in {item["id"] for item in store.list_items("source", "active")}
    restored = store.restore_item(source["id"])
    assert restored["state"] == "paused"

    other = next(item for item in store.list_items("source") if item["id"] != source["id"])
    with pytest.raises(ValueError, match="already exists"):
        store.update_item(other["id"], {"model": restored["model"]})


def test_prepare_legacy_filters_paused_and_trashed_items(tmp_path: Path) -> None:
    store = CatalogStore(tmp_path / "catalog.sqlite3")
    active = store.create_item("source", {"model": "55P7L", "source_sheets": ["TV"]})
    paused = store.create_item("source", {"model": "Q65H", "source_sheets": ["SB"], "paused": True})
    stock = store.create_item(
        "stock",
        {"nomenclature": "TCL 55P7L", "model": "55P7L", "warehouse": "Main", "quantity": 2, "unit_cost_eur": 300},
    )
    store.create_item(
        "stock",
        {"nomenclature": "Paused", "model": "Q65H", "quantity": 5, "unit_cost_eur": 100, "paused": True},
    )

    legacy_db = tmp_path / "legacy.sqlite3"
    active_book = tmp_path / "active.xlsx"
    store.prepare_legacy(tmp_path / "missing.sqlite3", legacy_db, active_book)

    from openpyxl import load_workbook
    import sqlite3

    workbook = load_workbook(active_book, data_only=True)
    assert [row[0].value for row in workbook["TV"].iter_rows(min_row=2)] == [active["model"]]
    assert [row[0].value for row in workbook["SB"].iter_rows(min_row=2)] == []
    with sqlite3.connect(legacy_db) as db:
        rows = db.execute("SELECT nomenclature, canonical_model FROM stock_items").fetchall()
    assert rows == [(stock["nomenclature"], "55P7L")]


def test_source_toggles_and_manual_shop_links_are_persisted(tmp_path: Path) -> None:
    store = CatalogStore(tmp_path / "catalog.sqlite3")
    item = store.create_item(
        "source",
        {"model": "55P7L", "shop_links": {"varle": "https://www.varle.lt/example.html"}},
    )
    assert item["shop_links"] == {"varle": "https://www.varle.lt/example.html"}

    store.update_item(item["id"], {"shop_links": {"elesen": "https://www.elesen.lt/example"}})
    assert store.get_item(item["id"])["shop_links"] == {"elesen": "https://www.elesen.lt/example"}

    store.update_source("senukai", False)
    store.update_source("senukai", collection_method="playwright")
    store.remember_source_method("senukai", "playwright")
    store.update_source_master("shop", False)
    senukai = next(source for source in store.list_sources() if source["key"] == "senukai")
    assert senukai["enabled"] is False
    assert senukai["master_enabled"] is False
    assert senukai["effective_enabled"] is False
    assert senukai["collection_method"] == "playwright"
    assert senukai["last_success_method"] == "playwright"


def test_discovered_links_and_monitoring_history_are_persisted(tmp_path: Path) -> None:
    store = CatalogStore(tmp_path / "catalog.sqlite3")
    item = store.create_item(
        "source", {"model": "65P7K", "shop_links": {"varle": "https://www.varle.lt/old"}}
    )
    store.remember_source_link(item["id"], "rde", "https://www.rde.ee/product/65p7k")
    store.remember_source_link(item["id"], "varle", "https://www.varle.lt/new")
    store.remember_source_link(item["id"], "salidzini", "https://www.salidzini.lv/cena?q=65P7K")

    links = store.get_item(item["id"])
    assert links["shop_links"] == {
        "varle": "https://www.varle.lt/new", "rde": "https://www.rde.ee/product/65p7k"
    }
    assert links["marketplace_links"]["salidzini"].endswith("q=65P7K")

    store.register_monitoring_session("run-history", False, ["salidzini"])
    history = store.list_monitoring_sessions()
    assert history[0]["run_id"] == "run-history"
    assert history[0]["status"] == "COMPLETE"


def test_stock_item_can_be_promoted_and_trash_can_be_deleted_permanently(tmp_path: Path) -> None:
    store = CatalogStore(tmp_path / "catalog.sqlite3")
    stock = store.create_item("stock", {"nomenclature": "TCL TV", "model": "55T7B"})

    promoted = store.promote_stock_item(stock["id"])
    assert promoted["kind"] == "source"
    assert promoted["model"] == "55T7B"
    assert promoted["paused"] is False
    assert store.promote_stock_item(stock["id"])["id"] == promoted["id"]

    inferred_stock = store.create_item("stock", {"nomenclature": "Monitor TCL 25G64, gab."})
    inferred = store.promote_stock_item(inferred_stock["id"])
    assert inferred["model"] == "25G64"
    assert inferred["source_sheets"] == ["Monitors"]
    assert store.get_item(inferred_stock["id"])["model"] == "25G64"
    assert store.get_item(inferred_stock["id"])["matched"] is True

    store.trash_item(promoted["id"])
    store.delete_item_permanently(promoted["id"])
    with pytest.raises(KeyError):
        store.get_item(promoted["id"])


def test_action_required_shop_observation_can_be_retried(tmp_path: Path) -> None:
    store = CatalogStore(tmp_path / "catalog.sqlite3")
    source = store.create_item("source", {"model": "25G64", "source_sheets": ["Monitors"]})
    shop = next(item for item in store.list_sources() if item["key"] == "bite")
    store.register_monitoring_session("run-1", False, [])
    store.start_shop_run("run-1", [source], [shop])
    store.finish_shop_observation(
        "run-1", source["id"], "bite", "ACTION_REQUIRED",
        search_url="https://www.bite.lt/paieska?q=25G64", error="Verification required",
    )

    store.retry_shop_observation("run-1", source["id"], "bite")

    result = store.shop_run("run-1")
    assert result["status"] == "RUNNING"
    assert result["results"][0]["status"] == "PENDING"
    assert result["results"][0]["error"] is None

    store.finish_shop_observation(
        "run-1", source["id"], "bite", "ACTION_REQUIRED", error="Verification required"
    )
    resolved = store.resolve_shop_observation("run-1", source["id"], "bite", "NOT_FOUND")
    assert resolved["status"] == "NOT_FOUND"
    assert resolved["collection_method"] == "manual"
    assert resolved["attempts"][0]["result"] == "CONFIRMED"


def test_manual_decision_history_and_legacy_marketplace_correction(tmp_path: Path) -> None:
    store = CatalogStore(tmp_path / "catalog.sqlite3")
    source = store.create_item("source", {"model": "25G64"})
    varle = next(item for item in store.list_sources() if item["key"] == "varle")
    store.register_monitoring_session("manual-old", False, ["hinnavaatlus"])
    store.start_shop_run("manual-old", [source], [varle])
    store.finish_shop_observation(
        "manual-old", source["id"], "varle", "ACTION_REQUIRED",
        search_url="https://www.varle.lt/search/?q=25G64",
    )
    store.resolve_shop_observation(
        "manual-old", source["id"], "varle", "SUCCESS", price_eur=199,
    )

    store.register_monitoring_session("manual-new", False, ["hinnavaatlus"])
    store.start_shop_run("manual-new", [source], [varle])
    previous = store.previous_manual_resolution(
        "manual-new", source["id"], "shop", "varle"
    )
    assert previous["status"] == "SUCCESS"
    assert previous["price_eur"] == 199
    assert previous["decided_at"]

    corrected = store.resolve_marketplace_result(
        "manual-new", source["id"], "hinnavaatlus", "NOT_FOUND"
    )
    assert corrected["status"] == "NOT_FOUND"
    assert store.manual_resolution(
        "manual-new", source["id"], "marketplace", "hinnavaatlus"
    )["status"] == "NOT_FOUND"


def test_hard_stop_finalizes_shop_and_legacy_work(tmp_path: Path) -> None:
    store = CatalogStore(tmp_path / "catalog.sqlite3")
    source = store.create_item("source", {"model": "55T7B"})
    shop = next(item for item in store.list_sources() if item["key"] == "varle")
    store.register_monitoring_session("run-stop", True, ["hinnavaatlus"])
    store.start_shop_run("run-stop", [source], [shop])

    assert store.cancel_shop_run("run-stop") == 1
    store.stop_monitoring_session("run-stop")
    assert store.shop_run("run-stop")["results"][0]["status"] == "INCOMPLETE"
    assert store.monitoring_session("run-stop")["stopped_at"] is not None

    legacy_database = tmp_path / "legacy.sqlite3"
    with sqlite3.connect(legacy_database) as db:
        db.executescript(
            "CREATE TABLE monitoring_runs(id TEXT PRIMARY KEY, status TEXT, finished_at TEXT, error TEXT);"
            "CREATE TABLE marketplace_tasks(run_id TEXT, status TEXT, finished_at TEXT, error TEXT);"
            "INSERT INTO monitoring_runs(id,status) VALUES('run-stop','RUNNING');"
            "INSERT INTO marketplace_tasks(run_id,status) VALUES('run-stop','RUNNING');"
        )
    assert store.cancel_legacy_run(legacy_database, "run-stop") == 1
    with sqlite3.connect(legacy_database) as db:
        assert db.execute("SELECT status FROM monitoring_runs").fetchone()[0] == "INCOMPLETE"
        assert db.execute("SELECT status FROM marketplace_tasks").fetchone()[0] == "INCOMPLETE"


def test_clear_monitoring_session_removes_current_results_and_stays_latest(tmp_path: Path) -> None:
    store = CatalogStore(tmp_path / "catalog.sqlite3")
    source = store.create_item("source", {"model": "25G64"})
    shop = next(item for item in store.list_sources() if item["key"] == "bite")
    salidzini = next(item for item in store.list_sources() if item["key"] == "salidzini")
    store.register_monitoring_session("run-clear", False, ["salidzini"])
    store.start_shop_run("run-clear", [source], [shop])
    store.start_assisted_marketplace_run("run-clear", [source], [salidzini])

    store.clear_monitoring_session("run-clear")

    session = store.monitoring_session("run-clear")
    assert session is not None and session["cleared_at"] is not None
    assert store.latest_monitoring_session()["run_id"] == "run-clear"
    assert store.shop_run("run-clear")["results"] == []
    assert store.assisted_marketplace_run("run-clear")["results"] == []


def test_shop_cache_and_protection_cooldown_persist_between_runs(tmp_path: Path) -> None:
    store = CatalogStore(tmp_path / "catalog.sqlite3")
    source = store.create_item("source", {"model": "25G64"})
    shop = next(item for item in store.list_sources() if item["key"] == "bite")

    store.register_monitoring_session("cache-source", False, [])
    store.start_shop_run("cache-source", [source], [shop])
    store.finish_shop_observation(
        "cache-source", source["id"], "bite", "SUCCESS", title="TCL 25G64",
        price_eur=176.4, availability="IN_STOCK", product_url="https://www.bite.lt/example",
    )
    store.register_monitoring_session("cache-target", False, [])
    store.start_shop_run("cache-target", [source], [shop], cache_ttl_seconds=14_400)
    cached = store.shop_run("cache-target")["results"][0]
    assert cached["status"] == "SUCCESS"
    assert cached["cached"] is True
    assert cached["price_eur"] == 176.4
    assert not store.shop_observation_pending("cache-target", source["id"], "bite")

    store.register_monitoring_session("blocked-source", False, [])
    store.start_shop_run("blocked-source", [source], [shop])
    retry_after = store.pause_shop_for_protection("blocked-source", "bite", 3600, "Rate limited")
    blocked = store.shop_run("blocked-source")["results"][0]
    assert blocked["status"] == "COOLDOWN"
    assert blocked["retry_after"] == retry_after

    store.register_monitoring_session("blocked-target", False, [])
    store.start_shop_run("blocked-target", [source], [shop], cache_ttl_seconds=0)
    assert store.shop_run("blocked-target")["results"][0]["status"] == "COOLDOWN"
    store.retry_shop_observation("blocked-target", source["id"], "bite")
    assert store.shop_run("blocked-target")["results"][0]["status"] == "PENDING"
