from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook

from price_monitor_v5.catalog import CatalogStore
from price_monitor_v5.exporter import write_monitoring_export
from price_monitor_v5.offers import parse_page, seller_key


@pytest.mark.parametrize("name", ["Smartech", "Smartech.ee", "Smartech Shop", "SMARTECH SHOP", "  Smartech   Shop  "])
def test_smartech_seller_aliases(name):
    assert seller_key(name, "hinnavaatlus") == "smartech"


@pytest.mark.parametrize("name", ["Smartech.lv", "Smartech.lt", "Other Smartech Shop", "Smartech Shop Outlet"])
def test_no_fuzzy_smartech_match(name):
    assert seller_key(name, "hinnavaatlus") is None


def test_country_guard_still_applies():
    assert seller_key("Smartech Shop", "salidzini") is None
    assert seller_key("Smartech Shop", "kaina24") is None
    assert seller_key("Smartech.ee", "salidzini") == "smartech"


def test_hinnavaatlus_smartech_shop_row():
    html = '''<h1>Teler TCL 32" 32S4K</h1><table><tr class="offer">
    <td><a class="name subtitle-main" href="/dealer/123-example">Smartech Shop</a></td>
    <td class="stock"><div class="in-stock"><div class="in-stock-message">2-4 p</div></div></td>
    <td><button class="offer-price">157,60 €</button><button class="offer-per-month">6,17 €/kuus</button></td>
    </tr></table>'''
    result = parse_page("hinnavaatlus", "32S4K", html, "https://www.hinnavaatlus.ee/6403045/tcl-32-32s4k/")
    assert len(result["offers"]) == 1
    offer = result["offers"][0]
    assert offer["store"] == "Smartech Shop"
    assert offer["seller_key"] == "smartech"
    assert offer["price_eur"] == 157.6
    assert offer["availability"] == "UNKNOWN"


def test_stored_and_cached_offer_remapped_without_rechecking_or_rewriting_history(tmp_path):
    store = CatalogStore(tmp_path / "catalog.db")
    item = store.create_item("source", {"model": "32S4K"})
    store.update_source("kaina24", False)
    store.update_source("salidzini", False)
    store.begin_run("old", "deep")
    task = store.check("old", item["id"], "hinnavaatlus")
    checked_at = (datetime.now(UTC) - timedelta(minutes=2)).isoformat()
    task.update(status="SUCCESS", coverage="complete", collection_method="marketplace HTML", finished_at=checked_at,
                offers=[{"store": "Smartech Shop", "seller_key": None, "price_eur": 157.6,
                         "title": 'Teler TCL 32" 32S4K', "availability": "UNKNOWN", "manual": False,
                         "marketplace_key": "hinnavaatlus", "url": "https://www.hinnavaatlus.ee/6403045/tcl-32-32s4k/"}])
    store.save_check(task)
    original = store.check("old", item["id"], "hinnavaatlus")
    store.begin_run("cached", "balanced")
    assert store.check("cached", item["id"], "hinnavaatlus")["cached"]
    for run_id in ("old", "cached"):
        run = store.run(run_id)
        shop = next(s for s in run["shop_results"] if s["shop_key"] == "smartech")
        assert shop["status"] == "SUCCESS"
        assert shop["price_eur"] == 157.6
        assert shop["checked_at"] == checked_at
        assert shop["observations"][0]["store"] == "Smartech Shop"
        assert shop["observations"][0]["cached"] == (run_id == "cached")
        assert run["tasks"][0]["cheapest_in_stock"] is None
        export_path = tmp_path / f"{run_id}.xlsx"
        write_monitoring_export(export_path, run)
        book = load_workbook(export_path)
        column = next(cell.column for cell in book.active[1] if cell.value == "Smartech")
        assert book.active.cell(2, column).value == 157.6
        book.close()
    assert store.check("old", item["id"], "hinnavaatlus") == original

