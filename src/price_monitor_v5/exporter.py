from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime, UTC
from .sources import SHOPS


def safe(value):
    text = str(value or "")
    return "'"+text if text.startswith(("=","+","-","@")) else text


def checked_time(value):
    return datetime.fromisoformat(value).astimezone(UTC).replace(tzinfo=None) if value else None


def write_monitoring_export(path, run):
    book = Workbook()
    sheet = book.active
    sheet.title = "Monitoring"
    sheet.append(["Model","Marketplace min price","Marketplace max price","Lowest in-stock","Lowest pre-order",*[s.name for s in SHOPS],"Stock quantity","Unit cost, EUR","Status"])
    details = book.create_sheet("All seller offers")
    details.append(["Model","Marketplace","Seller","Price, EUR","Availability","Marketplace link","Checked at (UTC)","Method","Coverage","Price basis","Matched model"])
    checks = book.create_sheet("Checks")
    checks.append(["Model","Marketplace","Status","Coverage","Error","Checked at (UTC)"])
    for item_id in dict.fromkeys(t["item_id"] for t in run["tasks"]):
        tasks = [t for t in run["tasks"] if t["item_id"]==item_id]
        model = tasks[0]["source_model"]
        summaries = []
        for field in ("cheapest_in_stock","highest_in_stock"):
            lines = []
            for task in tasks:
                offer = task.get(field) or task.get("lowest_reported" if field == "cheapest_in_stock" else "highest_reported")
                if offer:
                    note = "" if task.get(field) else " [reported price; " + offer["availability"].replace("_", " ").lower() + "]"
                    if offer.get("price_basis") == "loyalty":
                        note += " [loyalty price]"
                    lines.append(f"{task['marketplace']}: {offer['price_eur']:.2f} EUR ({offer['store']})" + note + (" [partial]" if task.get("coverage") != "complete" else ""))
                else:
                    lines.append(f"{task['marketplace']}: {task['status']} / no in-stock price")
            summaries.append("\n".join(lines))
        all_offers = [o for t in tasks for o in t.get("offers",[])]
        low = [min((o["price_eur"] for o in all_offers if o["availability"]==state),default=None) for state in ("IN_STOCK","PRE_ORDER")]
        shops = [next((s for s in run["shop_results"] if s["item_id"]==item_id and s["shop_key"]==shop.key),{}) for shop in SHOPS]
        values = [safe(model),*summaries,*low,*[s.get("price_eur") if s.get("price_eur") is not None else s.get("status","—") for s in shops],tasks[0].get("stock_quantity"),tasks[0].get("stock_unit_cost_eur")," / ".join(dict.fromkeys(t["status"] for t in tasks))]
        sheet.append(values)
        sheet.row_dimensions[sheet.max_row].height = max(54, 36*len(tasks))
        for col, result in enumerate(shops,6):
            if result.get("product_url"):
                sheet.cell(sheet.max_row,col).hyperlink = result["product_url"]
                sheet.cell(sheet.max_row,col).font = Font(color="176A49",underline="single")
        # Excel supports one hyperlink per cell, so every marketplace offer has
        # its own numeric row and link on the details sheet.
        for task in tasks:
            checks.append([safe(model),task["marketplace"],task["status"],task.get("coverage","partial"),safe(task.get("error")),checked_time(task.get("finished_at"))])
            for offer in task.get("offers",[]):
                details.append([safe(model),task["marketplace"],safe(offer["store"]),offer["price_eur"],offer["availability"],offer["url"],checked_time(task.get("finished_at")),task.get("collection_method"),task.get("coverage"),offer.get("price_basis", "not recorded"),safe(offer.get("matched_model") or model)])
                details.cell(details.max_row,6).hyperlink = offer["url"]
    for ws in book:
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        ws.print_title_rows = "1:1"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
        ws.sheet_view.showGridLines = False
        for c in ws[1]:
            c.fill = PatternFill("solid",fgColor="176A49")
            c.font = Font(color="FFFFFF",bold=True)
            c.alignment = Alignment(wrap_text=True,vertical="center")
        ws.row_dimensions[1].height = 32
        for index in range(1,ws.max_column+1):
            ws.column_dimensions[get_column_letter(index)].width = 19
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True,vertical="top")
                if isinstance(cell.value,float):
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right",vertical="top")
                if isinstance(cell.value,datetime):
                    cell.number_format = 'yyyy-mm-dd hh:mm'
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["C"].width = 42
    details.column_dimensions["F"].width = 56
    checks.column_dimensions["E"].width = 65
    checks.column_dimensions["F"].width = 24
    details.column_dimensions["G"].width = 24
    for i in range(2,checks.max_row+1):
        checks.row_dimensions[i].height = 42
    path.parent.mkdir(parents=True,exist_ok=True)
    book.save(path)
