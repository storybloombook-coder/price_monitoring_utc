from pathlib import Path
from datetime import UTC, datetime, timedelta
import json
import re
import sqlite3
from openpyxl import load_workbook
from price_monitor_v4.catalog import CatalogStore as BaseCatalog, canonicalize, number, utc_now, source_sheet_for_nomenclature
from .sources import SOURCE_BY_KEY, search_url, marketplace_url, SHOPS
from .offers import exact_model, compact, seller_key


def infer_model(text):
    # Only model-shaped TCL identifiers; multiple candidates need explicit review.
    candidates = re.findall(r"(?<![A-Z0-9])(?:\d{2,3}[A-Z]\d[A-Z0-9]*|[A-Z]\d{2}[A-Z][A-Z0-9]*)(?:[ -]PRO)?(?![A-Z0-9])", str(text).upper())
    candidates = list(dict.fromkeys(candidates))
    return candidates[0] if len(candidates) == 1 else None


class CatalogStore(BaseCatalog):
    def list_items(self, kind, scope="active", query=""):
        if kind not in {"source","stock"} or scope not in {"active","all","trash"}:
            raise ValueError("Invalid catalog scope")
        clauses, args = ["kind=?"], [kind]
        if scope != "all":
            clauses.append("deleted_at IS NULL" if scope=="active" else "deleted_at IS NOT NULL")
        if query.strip():
            clauses.append("UPPER(COALESCE(model,'') || ' ' || COALESCE(nomenclature,'') || ' ' || COALESCE(warehouse,'')) LIKE ?")
            args.append("%"+query.strip().upper()+"%")
        with self.connect() as db:
            rows = db.execute(f"SELECT * FROM catalog_items WHERE {' AND '.join(clauses)} ORDER BY COALESCE(model,nomenclature),id",args).fetchall()
            active = {compact(r[0]) for r in db.execute("SELECT canonical_model FROM catalog_items WHERE kind='source' AND deleted_at IS NULL AND paused=0")}
            links = {}
            for row in db.execute("SELECT * FROM item_marketplace_links"):
                links.setdefault(row["item_id"],{})[row["marketplace_key"]] = row["product_url"]
        result = []
        for row in rows:
            item = dict(row)
            item["paused"] = bool(item["paused"])
            item["source_sheets"] = json.loads(item["source_sheets"] or "[]")
            item["state"] = "trash" if item["deleted_at"] else "paused" if item["paused"] else "active"
            if kind=="source":
                item.update(shop_links={},marketplace_links=links.get(item["id"],{}))
            else:
                item["matched"] = bool(item["canonical_model"]) and compact(item["canonical_model"]) in active
            result.append(item)
        return result

    def migrate(self):
        super().migrate()
        with self.connect() as db:
            db.executescript("""
                CREATE TABLE IF NOT EXISTS marketplace_checks (
                    run_id TEXT NOT NULL REFERENCES monitoring_sessions(run_id),
                    item_id INTEGER NOT NULL, marketplace_key TEXT NOT NULL,
                    model TEXT NOT NULL, canonical_model TEXT NOT NULL,
                    status TEXT NOT NULL, data TEXT NOT NULL,
                    PRIMARY KEY(run_id,item_id,marketplace_key));
                CREATE INDEX IF NOT EXISTS v5_check_cache ON marketplace_checks(canonical_model,marketplace_key);
                CREATE TABLE IF NOT EXISTS v5_manual_history (
                    id INTEGER PRIMARY KEY, run_id TEXT NOT NULL, item_id INTEGER NOT NULL,
                    marketplace_key TEXT NOT NULL, decision TEXT NOT NULL, decided_at TEXT NOT NULL);
            """)

    def _replace_shop_links(self, db, item_id, links):
        if any(links.values()):
            raise ValueError("Direct shop links are disabled in v5; use marketplace links")

    def _replace_marketplace_links(self, db, item_id, links):
        for key, url in links.items():
            if url:
                marketplace_url(key, url)
        super()._replace_marketplace_links(db, item_id, links)

    def remember_source_link(self, item_id, source_key, product_url):
        return super().remember_source_link(item_id, source_key, marketplace_url(source_key, product_url))

    def create_item(self, kind, payload):
        if kind=="source" and any(compact(i["canonical_model"])==compact(canonicalize(str(payload.get("model") or ""))) for i in self.list_items("source")):
            raise ValueError("This model already exists in monitoring")
        item = super().create_item(kind, payload)
        if kind == "stock":
            self.enroll_stock()
            item = self.get_item(item["id"])
        return item

    def update_item(self, item_id, payload):
        item = super().update_item(item_id, payload)
        if item["kind"] == "stock":
            self.enroll_stock()
            item = self.get_item(item_id)
        return item

    def enroll_stock(self):
        inserted = 0
        now = utc_now()
        with self._lock, self.connect() as db:
            source = {compact(row["canonical_model"]): row["id"] for row in db.execute("SELECT * FROM catalog_items WHERE kind='source'")}
            for row in db.execute("SELECT * FROM catalog_items WHERE kind='stock' AND deleted_at IS NULL").fetchall():
                model = row["model"] or infer_model(row["nomenclature"])
                if not model:
                    continue
                canonical = canonicalize(model)
                db.execute("UPDATE catalog_items SET model=?,canonical_model=?,updated_at=? WHERE id=?", (model, canonical, now, row["id"]))
                if compact(canonical) in source:
                    # Import never resumes a deliberately paused or trashed SKU.
                    continue
                result = db.execute("INSERT INTO catalog_items(kind,model,canonical_model,source_sheets,origin,created_at,updated_at) VALUES('source',?,?,?,'stock',?,?)",
                                    (canonical, canonical, json.dumps([source_sheet_for_nomenclature(row["nomenclature"] or "")]), now, now))
                source[compact(canonical)] = result.lastrowid
                inserted += 1
        return inserted

    def clear_catalog(self, kind):
        if kind not in {"source", "stock"}:
            raise ValueError("Choose source or stock")
        now = utc_now()
        with self._lock, self.connect() as db:
            changed = db.execute("UPDATE catalog_items SET deleted_at=?,updated_at=? WHERE kind=? AND deleted_at IS NULL", (now, now, kind)).rowcount
        return {"kind": kind, "moved_to_trash": changed}

    def import_stock_workbook(self, path, filename):
        parsed = []
        models = [i["model"] for i in self.list_items("source", "all")]
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            for sheet in workbook:
                warehouse = None
                for row in sheet.iter_rows(values_only=True):
                    name = str(row[1] or "").strip() if len(row) > 1 else ""
                    qty, cost = (row[2] if len(row)>2 else None), (row[3] if len(row)>3 else None)
                    if not name or name.lower() in {"item", "total", "nomenclature", "наименование"}:
                        continue
                    if qty in (None, "") and cost in (None, ""):
                        warehouse = name
                        continue
                    matches = [m for m in models if exact_model(m, name)]
                    model = matches[0] if len(matches) == 1 else infer_model(name) if not matches else None
                    parsed.append((name, warehouse, number(qty), number(cost), model, canonicalize(model) if model else None))
        finally:
            workbook.close()
        if not parsed:
            raise ValueError("No stock rows found in columns B-D")
        with self._lock, self.connect() as db:
            db.execute("DELETE FROM catalog_items WHERE kind='stock' AND origin IN ('file','v3')")
            for item in parsed:
                db.execute("INSERT INTO catalog_items(kind,nomenclature,warehouse,quantity,unit_cost_eur,model,canonical_model,source_filename,origin,created_at,updated_at) VALUES('stock',?,?,?,?,?,?,?,'file',?,?)", (*item, filename, utc_now(), utc_now()))
            self._set_meta(db, "stock_upload", {"filename": filename, "uploaded_at": utc_now(), "count": len(parsed)})
        inserted = self.enroll_stock()
        return {"imported": len(parsed), "matched": sum(bool(r[4]) for r in parsed), "enrolled": inserted, "needs_review": sum(not r[4] for r in parsed)}

    def begin_run(self, run_id, mode):
        items = {compact(i["canonical_model"]): i for i in self.list_items("source") if not i["paused"]}
        sources = [s for s in self.list_sources() if s["kind"] == "marketplace" and s["effective_enabled"]]
        if not items or not sources:
            raise ValueError("Add an active model and enable at least one marketplace")
        self.register_monitoring_session(run_id, False, [s["key"] for s in sources], mode)
        stock = self.list_items("stock")
        for item in items.values():
            matching_stock = [s for s in stock if compact(s.get("canonical_model")) == compact(item["canonical_model"])]
            quantity = sum(s["quantity"] or 0 for s in matching_stock)
            costs = [s["unit_cost_eur"] for s in matching_stock if s["unit_cost_eur"] is not None]
            cost = sum((s["quantity"] or 0)*(s["unit_cost_eur"] or 0) for s in matching_stock)/quantity if quantity else sum(costs)/len(costs) if costs else None
            for source in sources:
                key = source["key"]
                task = {"run_id": run_id, "item_id": item["id"], "source_model": item["model"], "canonical_model": item["canonical_model"],
                        "marketplace": source["name"], "marketplace_key": key, "status": "PENDING", "assisted": True,
                        "offers": [], "search_url": search_url(key, item["model"]), "product_url": item.get("marketplace_links", {}).get(key),
                        "stock_quantity": quantity, "stock_unit_cost_eur": cost, "cached": False,
                        "previous_manual_resolution": self.previous_decision(item["id"], key, run_id)}
                ttl = {"quick": 12, "balanced": 4, "deep": 0}[mode]
                cached = self.cached_check(item["canonical_model"], key, ttl)
                if cached:
                    task.update({k: cached.get(k) for k in ("offers", "status", "finished_at", "coverage", "collection_method", "product_url", "error")})
                    task["cached"] = True
                self.save_check(task)

    def cached_check(self, canonical, key, hours):
        if not hours:
            return None
        with self.connect() as db:
            rows = db.execute("SELECT c.data FROM marketplace_checks c JOIN monitoring_sessions s ON s.run_id=c.run_id WHERE c.canonical_model=? AND c.marketplace_key=? AND s.cleared_at IS NULL ORDER BY s.created_at DESC LIMIT 1", (canonical,key)).fetchall()
        threshold = (datetime.now(UTC)-timedelta(hours=hours)).isoformat()
        return next((t for r in rows if (t:=json.loads(r[0])).get("finished_at", "") >= threshold and t.get("status")=="SUCCESS" and t.get("coverage") == "complete" and t.get("collection_method") != "manual"), None)

    def save_check(self, task):
        with self._lock, self.connect() as db:
            db.execute("INSERT OR REPLACE INTO marketplace_checks VALUES(?,?,?,?,?,?,?)", (task["run_id"],task["item_id"],task["marketplace_key"],task["source_model"],task["canonical_model"],task["status"],json.dumps(task)))

    def checks(self, run_id):
        with self.connect() as db:
            return [json.loads(r[0]) for r in db.execute("SELECT data FROM marketplace_checks WHERE run_id=? ORDER BY model,marketplace_key", (run_id,))]

    def check(self, run_id, item_id, key):
        with self.connect() as db:
            row = db.execute("SELECT data FROM marketplace_checks WHERE run_id=? AND item_id=? AND marketplace_key=?", (run_id,item_id,key)).fetchone()
        if not row:
            raise KeyError("Check not found")
        return json.loads(row[0])

    def previous_decision(self, item_id, key, run_id):
        with self.connect() as db:
            row = db.execute("SELECT decision FROM v5_manual_history WHERE item_id=? AND marketplace_key=? AND run_id<>? ORDER BY id DESC LIMIT 1", (item_id,key,run_id)).fetchone()
        return json.loads(row[0]) if row else None

    def record_decision(self, task, decision):
        decision["decided_at"] = utc_now()
        with self.connect() as db:
            db.execute("INSERT INTO v5_manual_history(run_id,item_id,marketplace_key,decision,decided_at) VALUES(?,?,?,?,?)", (task["run_id"],task["item_id"],task["marketplace_key"],json.dumps(decision),decision["decided_at"]))

    def run(self, run_id):
        session = self.monitoring_session(run_id)
        if not session:
            raise KeyError("Run not found")
        tasks = [] if session.get("cleared_at") else self.checks(run_id)
        for task in tasks:
            # Seller identity is derived, not a price observation. Apply current
            # aliases to stored/cache results too, without rewriting prices,
            # timestamps, original seller labels or historical check records.
            offers = [{**offer, "seller_key": seller_key(offer.get("store"), task["marketplace_key"])}
                      for offer in task.get("offers", [])]
            task["offers"] = offers
            def lowest(state):
                return min((o for o in offers if o["availability"] == state), key=lambda o:o["price_eur"], default=None)
            task["cheapest_in_stock"] = lowest("IN_STOCK")
            task["cheapest_pre_order"] = lowest("PRE_ORDER")
            task["highest_in_stock"] = max((o for o in offers if o["availability"] == "IN_STOCK"),key=lambda o:o["price_eur"],default=None)
            # Report unconfirmed prices separately; never relabel them as in stock.
            task["lowest_reported"] = min(offers,key=lambda o:o["price_eur"],default=None)
            task["highest_reported"] = max(offers,key=lambda o:o["price_eur"],default=None)
        shops = []
        for item_id in dict.fromkeys(t["item_id"] for t in tasks):
            model_tasks = [t for t in tasks if t["item_id"] == item_id]
            for shop in SHOPS:
                observations = [{**o, "marketplace": t["marketplace"], "checked_at": t.get("finished_at"), "cached": t.get("cached",False)} for t in model_tasks for o in t.get("offers",[]) if o.get("seller_key") == shop.key]
                in_stock = [o for o in observations if o["availability"] == "IN_STOCK"]
                best = min(in_stock or observations, key=lambda o:o["price_eur"], default={})
                complete = all(t["status"] in {"SUCCESS","NOT_FOUND"} and t.get("coverage") == "complete" for t in model_tasks)
                status = "SUCCESS" if best else "NOT_LISTED" if complete else "UNVERIFIED"
                shops.append({"item_id": item_id, "model": model_tasks[0]["source_model"], "shop_key": shop.key,"shop_name": shop.name,"status":status,
                              "observations": observations, "price_eur": best.get("price_eur"), "availability": best.get("availability"),
                              "product_url":best.get("url"),"checked_at":best.get("checked_at"),"coverage":"complete" if complete else "partial"})
        pending = sum(t["status"] in {"PENDING","RUNNING"} for t in tasks)
        return {"run_id":run_id,"status":"RUNNING" if pending else "INCOMPLETE" if session.get("stopped_at") else "COMPLETE",
                "cleared": bool(session.get("cleared_at")),"created_at":session["created_at"],"run_mode":session["run_mode"],"tasks":tasks,"shop_results":shops,
                "execution":{"total":len(tasks),"finished":len(tasks)-pending,"phase_label":"Marketplace-only checks" if pending else "Marketplace checks finished"}}

    def list_monitoring_sessions(self, limit=20):
        with self.connect() as db:
            sessions = [dict(r) for r in db.execute("SELECT s.*, COUNT(c.item_id) AS assisted_checks, SUM(c.status IN ('PENDING','RUNNING')) AS pending FROM monitoring_sessions s LEFT JOIN marketplace_checks c ON c.run_id=s.run_id GROUP BY s.run_id ORDER BY s.created_at DESC LIMIT ?", (max(0,min(limit,100)),))]
        for session in sessions:
            session.update(status="CLEARED" if session.get("cleared_at") else "RUNNING" if session["pending"] else "INCOMPLETE" if session.get("stopped_at") else "COMPLETE",shop_checks=0)
        return sessions

    def import_v4_catalog(self, path):
        """Explicit one-time migration: catalog only, never direct-shop caches/history."""
        if self.list_items("source", "all") or self.list_items("stock", "all"):
            raise ValueError("v5 catalog must be empty before migration")
        original = sqlite3.connect(f"{Path(path).resolve().as_uri()}?mode=ro",uri=True)
        original.row_factory = sqlite3.Row
        try:
            rows = original.execute("SELECT * FROM catalog_items").fetchall()
            with self.connect() as db:
                for row in rows:
                    keys = list(row.keys())
                    db.execute(f"INSERT INTO catalog_items({','.join(keys)}) VALUES({','.join('?' for _ in keys)})",tuple(row))
                self._set_meta(db,"v4_migration",{"source":str(path),"at":utc_now(),"items":len(rows)})
        finally:
            original.close()
        self.enroll_stock()
        return len(rows)
