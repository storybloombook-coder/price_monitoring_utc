from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .sources import SHOPS


TERMINAL_PRIORITY = {
    "ACTION_REQUIRED": 7, "RUNNING": 6, "PENDING": 6, "FAILED": 5, "INCOMPLETE": 4,
    "SUCCESS": 3, "NOT_FOUND": 2, "NOT_STARTED": 1,
}


def price(offer: dict[str, Any] | None) -> float | None:
    if not offer or offer.get("price_eur") in (None, ""):
        return None
    try:
        return float(offer["price_eur"])
    except (TypeError, ValueError):
        return None


def lowest(tasks: list[dict[str, Any]], field: str) -> dict[str, Any] | None:
    offers = [task.get(field) for task in tasks if task.get(field)]
    offers = [offer for offer in offers if price(offer) is not None]
    return min(offers, key=lambda offer: price(offer) or 0, default=None)


def lowest_with_shops(
    tasks: list[dict[str, Any]], shops: list[dict[str, Any]], field: str, availability: str
) -> dict[str, Any] | None:
    offers = [task.get(field) for task in tasks if task.get(field)]
    offers.extend(
        {
            "price_eur": item.get("price_eur"),
            "store": item.get("shop_name") or item.get("shop_key"),
            "url": item.get("product_url") or item.get("search_url"),
        }
        for item in shops
        if item.get("status") == "SUCCESS" and item.get("availability") == availability
    )
    valid = [offer for offer in offers if price(offer) is not None]
    return min(valid, key=lambda offer: price(offer) or 0, default=None)


def overall_status(tasks: list[dict[str, Any]], shops: list[dict[str, Any]]) -> str:
    statuses = [str(item.get("status") or "NOT_STARTED") for item in [*tasks, *shops]]
    return max(statuses, key=lambda value: TERMINAL_PRIORITY.get(value, 0), default="NOT_STARTED")


def autosize(sheet) -> None:
    for column in sheet.columns:
        letter = column[0].column_letter
        width = max((len(str(cell.value or "")) for cell in column), default=0)
        sheet.column_dimensions[letter].width = min(max(width + 2, 11), 42)


def write_monitoring_export(
    path: Path,
    run: dict[str, Any],
    models: list[dict[str, Any]],
) -> Path:
    tasks_by_model: dict[str, list[dict[str, Any]]] = defaultdict(list)
    shops_by_model: dict[str, list[dict[str, Any]]] = defaultdict(list)
    model_names: dict[str, str] = {}
    for item in models:
        key = str(item.get("canonical_model") or item.get("model") or "").upper()
        model_names[key] = str(item.get("model") or key)
    for task in run.get("tasks", []):
        key = str(task.get("canonical_model") or task.get("source_model") or "").upper()
        tasks_by_model[key].append(task)
        model_names.setdefault(key, str(task.get("source_model") or key))
    for item in run.get("shop_results", []):
        key = str(item.get("model") or "").upper()
        shops_by_model[key].append(item)
        model_names.setdefault(key, str(item.get("model") or key))

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Monitoring summary"
    headers = [
        "Model", "Marketplaces", "Lowest in-stock", "Lowest pre-order",
        *[shop.name for shop in SHOPS], "Stock quantity", "Stock unit cost EUR", "Margin EUR", "Status",
    ]
    summary.append(headers)
    shop_index = {shop.key: shop for shop in SHOPS}
    for key in sorted(model_names, key=lambda value: model_names[value].upper()):
        tasks = tasks_by_model[key]
        shop_results = shops_by_model[key]
        stock_offer = lowest_with_shops(tasks, shop_results, "cheapest_in_stock", "IN_STOCK")
        preorder_offer = lowest_with_shops(tasks, shop_results, "cheapest_pre_order", "PRE_ORDER")
        best_price = price(stock_offer) if stock_offer else price(preorder_offer)
        stock_quantity = next((task.get("stock_quantity") for task in tasks if task.get("stock_quantity") is not None), None)
        stock_cost = next((task.get("stock_unit_cost_eur") for task in tasks if task.get("stock_unit_cost_eur") is not None), None)
        try:
            margin = best_price - float(stock_cost) if best_price is not None and stock_cost is not None else None
        except (TypeError, ValueError):
            margin = None
        shop_values = {
            item["shop_key"]: price(item) if item.get("status") == "SUCCESS" else item.get("status")
            for item in shop_results if item.get("shop_key") in shop_index
        }
        summary.append([
            model_names[key],
            ", ".join(f"{task.get('marketplace')}: {task.get('status')}" for task in tasks),
            price(stock_offer), price(preorder_offer),
            *[shop_values.get(shop.key) for shop in SHOPS],
            stock_quantity, stock_cost, margin, overall_status(tasks, shop_results),
        ])

    details = workbook.create_sheet("Offer details")
    details.append(["Model", "Source type", "Source", "Status", "Title", "Price EUR", "Availability", "URL", "Checked at", "Error"])
    for key, tasks in tasks_by_model.items():
        for task in tasks:
            for field, availability in (("cheapest_in_stock", "IN_STOCK"), ("cheapest_pre_order", "PRE_ORDER")):
                offer = task.get(field)
                if offer:
                    details.append([
                        model_names[key], "Marketplace", task.get("marketplace"), task.get("status"),
                        task.get("matched_title"), price(offer), availability, offer.get("url"),
                        task.get("finished_at"), task.get("error"),
                    ])
            if not task.get("cheapest_in_stock") and not task.get("cheapest_pre_order"):
                details.append([model_names[key], "Marketplace", task.get("marketplace"), task.get("status"), task.get("matched_title"), None, None, task.get("product_url"), task.get("finished_at"), task.get("error")])
    for key, items in shops_by_model.items():
        for item in items:
            details.append([
                model_names[key], "Shop", item.get("shop_name") or item.get("shop_key"), item.get("status"),
                item.get("title"), price(item), item.get("availability"),
                item.get("product_url") or item.get("search_url"), item.get("checked_at"), item.get("error"),
            ])

    header_fill = PatternFill("solid", fgColor="146C43")
    for sheet in (summary, details):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        autosize(sheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
