from __future__ import annotations

import json
import re
import shutil
import sqlite3
import threading
import unicodedata
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook

from .sources import SOURCE_BY_KEY, SOURCES


SOURCE_SHEETS = ("TV", "SB", "Monitors")


def utc_now() -> str:
    return datetime.now(UTC).isoformat()


def canonicalize(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", value or "").upper().strip()
    return re.sub(r"\s+", " ", normalized)


def number(value: Any, default: float = 0.0) -> float:
    if value in (None, ""):
        return default
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = str(value).replace(" ", "").replace(",", ".")
    try:
        return float(cleaned)
    except ValueError:
        return default


def infer_model_from_nomenclature(value: str) -> str | None:
    """Extract a likely TCL model token such as 55T7B, Q65H, or 25G64."""
    tokens = re.findall(r"[A-Z0-9]+", canonicalize(value))
    ignored = {"TCL", "TV", "MONITOR", "SOUNDBAR", "GAB", "PRO"}
    for index, token in enumerate(tokens):
        if token in ignored or len(token) < 4:
            continue
        if any(char.isalpha() for char in token) and any(char.isdigit() for char in token):
            return f"{token} PRO" if index + 1 < len(tokens) and tokens[index + 1] == "PRO" else token
    return None


def source_sheet_for_nomenclature(value: str) -> str:
    lowered = value.lower()
    if "soundbar" in lowered:
        return "SB"
    if "monitor" in lowered:
        return "Monitors"
    return "TV"


class CatalogStore:
    def __init__(self, database: Path):
        self.database = database
        self.database.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        self.migrate()

    def connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.database, timeout=30)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute("PRAGMA journal_mode = WAL")
        return connection

    def migrate(self) -> None:
        with self.connect() as db:
            db.executescript(
                """
                CREATE TABLE IF NOT EXISTS catalog_items (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    kind TEXT NOT NULL CHECK(kind IN ('source', 'stock')),
                    model TEXT,
                    canonical_model TEXT,
                    source_sheets TEXT NOT NULL DEFAULT '[]',
                    nomenclature TEXT,
                    warehouse TEXT,
                    quantity REAL,
                    unit_cost_eur REAL,
                    source_filename TEXT,
                    origin TEXT NOT NULL DEFAULT 'manual',
                    paused INTEGER NOT NULL DEFAULT 0,
                    deleted_at TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_catalog_kind_state
                    ON catalog_items(kind, deleted_at, paused);
                CREATE INDEX IF NOT EXISTS idx_catalog_canonical
                    ON catalog_items(canonical_model);
                CREATE TABLE IF NOT EXISTS catalog_meta (
                    key TEXT PRIMARY KEY,
                    value TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS monitoring_sources (
                    key TEXT PRIMARY KEY,
                    name TEXT NOT NULL,
                    kind TEXT NOT NULL CHECK(kind IN ('marketplace', 'shop')),
                    country TEXT NOT NULL,
                    base_url TEXT NOT NULL,
                    enabled INTEGER NOT NULL DEFAULT 1,
                    collection_method TEXT NOT NULL DEFAULT 'auto',
                    last_success_method TEXT,
                    sort_order INTEGER NOT NULL
                );
                CREATE TABLE IF NOT EXISTS item_shop_links (
                    item_id INTEGER NOT NULL REFERENCES catalog_items(id) ON DELETE CASCADE,
                    shop_key TEXT NOT NULL REFERENCES monitoring_sources(key),
                    product_url TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY(item_id, shop_key)
                );
                CREATE TABLE IF NOT EXISTS item_marketplace_links (
                    item_id INTEGER NOT NULL REFERENCES catalog_items(id) ON DELETE CASCADE,
                    marketplace_key TEXT NOT NULL REFERENCES monitoring_sources(key),
                    product_url TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY(item_id, marketplace_key)
                );
                CREATE TABLE IF NOT EXISTS monitoring_sessions (
                    run_id TEXT PRIMARY KEY,
                    legacy_started INTEGER NOT NULL,
                    marketplace_keys TEXT NOT NULL DEFAULT '[]',
                    created_at TEXT NOT NULL,
                    stopped_at TEXT,
                    cleared_at TEXT
                );
                CREATE TABLE IF NOT EXISTS shop_observations (
                    run_id TEXT NOT NULL REFERENCES monitoring_sessions(run_id) ON DELETE CASCADE,
                    item_id INTEGER NOT NULL REFERENCES catalog_items(id),
                    shop_key TEXT NOT NULL REFERENCES monitoring_sources(key),
                    status TEXT NOT NULL,
                    title TEXT,
                    price_eur REAL,
                    availability TEXT,
                    product_url TEXT,
                    search_url TEXT,
                    error TEXT,
                    checked_at TEXT,
                    retry_after TEXT,
                    cached INTEGER NOT NULL DEFAULT 0,
                    collection_method TEXT,
                    attempts_json TEXT NOT NULL DEFAULT '[]',
                    PRIMARY KEY(run_id, item_id, shop_key)
                );
                CREATE INDEX IF NOT EXISTS idx_shop_observations_run ON shop_observations(run_id, status);
                CREATE TABLE IF NOT EXISTS assisted_marketplace_observations (
                    run_id TEXT NOT NULL REFERENCES monitoring_sessions(run_id) ON DELETE CASCADE,
                    item_id INTEGER NOT NULL REFERENCES catalog_items(id),
                    marketplace_key TEXT NOT NULL REFERENCES monitoring_sources(key),
                    status TEXT NOT NULL,
                    title TEXT,
                    seller_name TEXT,
                    price_eur REAL,
                    availability TEXT,
                    product_url TEXT,
                    search_url TEXT,
                    error TEXT,
                    checked_at TEXT,
                    cached INTEGER NOT NULL DEFAULT 0,
                    collection_method TEXT,
                    attempts_json TEXT NOT NULL DEFAULT '[]',
                    PRIMARY KEY(run_id, item_id, marketplace_key)
                );
                CREATE INDEX IF NOT EXISTS idx_assisted_marketplace_run
                    ON assisted_marketplace_observations(run_id, status);
                CREATE TABLE IF NOT EXISTS shop_protection_state (
                    shop_key TEXT PRIMARY KEY REFERENCES monitoring_sources(key),
                    cooldown_until TEXT NOT NULL,
                    reason TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                """
            )
            session_columns = {
                row[1] for row in db.execute("PRAGMA table_info(monitoring_sessions)")
            }
            if "stopped_at" not in session_columns:
                db.execute("ALTER TABLE monitoring_sessions ADD COLUMN stopped_at TEXT")
            if "cleared_at" not in session_columns:
                db.execute("ALTER TABLE monitoring_sessions ADD COLUMN cleared_at TEXT")
            observation_columns = {
                row[1] for row in db.execute("PRAGMA table_info(shop_observations)")
            }
            if "retry_after" not in observation_columns:
                db.execute("ALTER TABLE shop_observations ADD COLUMN retry_after TEXT")
            if "cached" not in observation_columns:
                db.execute("ALTER TABLE shop_observations ADD COLUMN cached INTEGER NOT NULL DEFAULT 0")
            if "collection_method" not in observation_columns:
                db.execute("ALTER TABLE shop_observations ADD COLUMN collection_method TEXT")
            if "attempts_json" not in observation_columns:
                db.execute("ALTER TABLE shop_observations ADD COLUMN attempts_json TEXT NOT NULL DEFAULT '[]'")
            source_columns = {
                row[1] for row in db.execute("PRAGMA table_info(monitoring_sources)")
            }
            if "collection_method" not in source_columns:
                db.execute("ALTER TABLE monitoring_sources ADD COLUMN collection_method TEXT NOT NULL DEFAULT 'auto'")
            if "last_success_method" not in source_columns:
                db.execute("ALTER TABLE monitoring_sources ADD COLUMN last_success_method TEXT")
            for order, source in enumerate(SOURCES):
                db.execute(
                    "INSERT INTO monitoring_sources(key, name, kind, country, base_url, enabled, sort_order) "
                    "VALUES(?, ?, ?, ?, ?, 1, ?) ON CONFLICT(key) DO UPDATE SET name=excluded.name, "
                    "kind=excluded.kind, country=excluded.country, base_url=excluded.base_url, sort_order=excluded.sort_order",
                    (source.key, source.name, source.kind, source.country, source.base_url, order),
                )

    def _set_meta(self, db: sqlite3.Connection, key: str, value: Any) -> None:
        db.execute(
            "INSERT INTO catalog_meta(key, value) VALUES(?, ?) "
            "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (key, json.dumps(value, ensure_ascii=False)),
        )

    def get_meta(self, key: str, default: Any = None) -> Any:
        with self.connect() as db:
            row = db.execute("SELECT value FROM catalog_meta WHERE key=?", (key,)).fetchone()
        return json.loads(row["value"]) if row else default

    def seed_from_v3(self, legacy_database: Path, source_workbook: Path | None = None) -> dict[str, int]:
        with self.connect() as db:
            existing = db.execute("SELECT COUNT(*) FROM catalog_items").fetchone()[0]
        if existing:
            return {"source": 0, "stock": 0}

        source_count = 0
        stock_count = 0
        if legacy_database.exists():
            legacy = sqlite3.connect(legacy_database)
            legacy.row_factory = sqlite3.Row
            try:
                tables = {row[0] for row in legacy.execute("SELECT name FROM sqlite_master WHERE type='table'")}
                with self.connect() as db:
                    now = utc_now()
                    if "products" in tables:
                        for row in legacy.execute("SELECT canonical_model, model, source_sheets FROM products"):
                            sheets = row["source_sheets"] or "[]"
                            db.execute(
                                "INSERT INTO catalog_items(kind, model, canonical_model, source_sheets, origin, created_at, updated_at) "
                                "VALUES('source', ?, ?, ?, 'v3', ?, ?)",
                                (row["model"], row["canonical_model"], sheets, now, now),
                            )
                            source_count += 1
                    if "stock_items" in tables:
                        for row in legacy.execute(
                            "SELECT canonical_model, nomenclature, warehouse, quantity, unit_cost_eur, source_filename FROM stock_items"
                        ):
                            db.execute(
                                "INSERT INTO catalog_items(kind, model, canonical_model, nomenclature, warehouse, quantity, "
                                "unit_cost_eur, source_filename, origin, created_at, updated_at) "
                                "VALUES('stock', ?, ?, ?, ?, ?, ?, ?, 'v3', ?, ?)",
                                (
                                    row["canonical_model"], row["canonical_model"], row["nomenclature"],
                                    row["warehouse"], number(row["quantity"]), number(row["unit_cost_eur"]),
                                    row["source_filename"], now, now,
                                ),
                            )
                            stock_count += 1
            finally:
                legacy.close()

        if source_count == 0 and source_workbook and source_workbook.exists():
            source_count = self.import_source_workbook(source_workbook, source_workbook.name)["imported"]
        return {"source": source_count, "stock": stock_count}

    def _row_to_item(self, row: sqlite3.Row, active_models: set[str]) -> dict[str, Any]:
        item = dict(row)
        item["paused"] = bool(item["paused"])
        item["source_sheets"] = json.loads(item["source_sheets"] or "[]")
        item["state"] = "trash" if item["deleted_at"] else ("paused" if item["paused"] else "active")
        if item["kind"] == "stock":
            item["matched"] = bool(item["canonical_model"] and item["canonical_model"] in active_models)
        else:
            with self.connect() as db:
                item["shop_links"] = {
                    result["shop_key"]: result["product_url"]
                    for result in db.execute(
                        "SELECT shop_key, product_url FROM item_shop_links WHERE item_id=?", (item["id"],)
                    )
                }
                item["marketplace_links"] = {
                    result["marketplace_key"]: result["product_url"]
                    for result in db.execute(
                        "SELECT marketplace_key, product_url FROM item_marketplace_links WHERE item_id=?",
                        (item["id"],),
                    )
                }
        return item

    def _replace_shop_links(self, db: sqlite3.Connection, item_id: int, links: dict[str, Any]) -> None:
        valid_shops = {source.key for source in SOURCES if source.kind == "shop"}
        db.execute("DELETE FROM item_shop_links WHERE item_id=?", (item_id,))
        now = utc_now()
        for key, value in links.items():
            url = str(value or "").strip()
            if key in valid_shops and url:
                db.execute(
                    "INSERT INTO item_shop_links(item_id, shop_key, product_url, updated_at) VALUES(?, ?, ?, ?)",
                    (item_id, key, url, now),
                )

    def _replace_marketplace_links(self, db: sqlite3.Connection, item_id: int, links: dict[str, Any]) -> None:
        valid_marketplaces = {source.key for source in SOURCES if source.kind == "marketplace"}
        db.execute("DELETE FROM item_marketplace_links WHERE item_id=?", (item_id,))
        now = utc_now()
        for key, value in links.items():
            url = str(value or "").strip()
            if key in valid_marketplaces and url:
                db.execute(
                    "INSERT INTO item_marketplace_links(item_id, marketplace_key, product_url, updated_at) "
                    "VALUES(?, ?, ?, ?)",
                    (item_id, key, url, now),
                )

    def remember_source_link(self, item_id: int, source_key: str, product_url: str) -> None:
        """Remember a verified product URL without replacing links for other sources."""
        source = SOURCE_BY_KEY.get(source_key)
        url = str(product_url or "").strip()
        if not source or not url:
            raise ValueError("A known source and product URL are required")
        with self._lock, self.connect() as db:
            item = db.execute(
                "SELECT kind FROM catalog_items WHERE id=? AND deleted_at IS NULL", (item_id,)
            ).fetchone()
            if not item or item["kind"] != "source":
                raise KeyError(item_id)
            if source.kind == "shop":
                db.execute(
                    "INSERT INTO item_shop_links(item_id,shop_key,product_url,updated_at) VALUES(?,?,?,?) "
                    "ON CONFLICT(item_id,shop_key) DO UPDATE SET product_url=excluded.product_url,"
                    "updated_at=excluded.updated_at",
                    (item_id, source_key, url, utc_now()),
                )
            else:
                db.execute(
                    "INSERT INTO item_marketplace_links(item_id,marketplace_key,product_url,updated_at) "
                    "VALUES(?,?,?,?) ON CONFLICT(item_id,marketplace_key) DO UPDATE SET "
                    "product_url=excluded.product_url,updated_at=excluded.updated_at",
                    (item_id, source_key, url, utc_now()),
                )

    def list_items(self, kind: str, scope: str = "active", query: str = "") -> list[dict[str, Any]]:
        if kind not in {"source", "stock"}:
            raise ValueError("kind must be source or stock")
        clauses = ["kind = ?"]
        values: list[Any] = [kind]
        if scope == "active":
            clauses.append("deleted_at IS NULL")
        elif scope == "trash":
            clauses.append("deleted_at IS NOT NULL")
        elif scope != "all":
            raise ValueError("scope must be active, trash or all")
        if query.strip():
            clauses.append("UPPER(COALESCE(model,'') || ' ' || COALESCE(nomenclature,'') || ' ' || COALESCE(warehouse,'')) LIKE ?")
            values.append(f"%{query.strip().upper()}%")

        with self.connect() as db:
            active_models = {
                row[0] for row in db.execute(
                    "SELECT canonical_model FROM catalog_items WHERE kind='source' AND deleted_at IS NULL AND paused=0"
                )
            }
            rows = db.execute(
                f"SELECT * FROM catalog_items WHERE {' AND '.join(clauses)} ORDER BY deleted_at IS NOT NULL, paused, COALESCE(model,nomenclature)",
                values,
            ).fetchall()
        return [self._row_to_item(row, active_models) for row in rows]

    def get_item(self, item_id: int) -> dict[str, Any]:
        with self.connect() as db:
            row = db.execute("SELECT * FROM catalog_items WHERE id=?", (item_id,)).fetchone()
            active_models = {
                result[0] for result in db.execute(
                    "SELECT canonical_model FROM catalog_items WHERE kind='source' AND deleted_at IS NULL AND paused=0"
                )
            }
        if not row:
            raise KeyError(item_id)
        return self._row_to_item(row, active_models)

    def create_item(self, kind: str, payload: dict[str, Any]) -> dict[str, Any]:
        now = utc_now()
        if kind == "source":
            model = str(payload.get("model", "")).strip()
            if not model:
                raise ValueError("Model is required")
            canonical = canonicalize(model)
            sheets = payload.get("source_sheets") or ["TV"]
            sheets = [sheet for sheet in sheets if sheet in SOURCE_SHEETS] or ["TV"]
            with self._lock, self.connect() as db:
                duplicate = db.execute(
                    "SELECT id FROM catalog_items WHERE kind='source' AND canonical_model=? AND deleted_at IS NULL",
                    (canonical,),
                ).fetchone()
                if duplicate:
                    raise ValueError(f"Model already exists (item {duplicate['id']})")
                cursor = db.execute(
                    "INSERT INTO catalog_items(kind, model, canonical_model, source_sheets, origin, paused, created_at, updated_at) "
                    "VALUES('source', ?, ?, ?, 'manual', ?, ?, ?)",
                    (model, canonical, json.dumps(sheets), int(bool(payload.get("paused"))), now, now),
                )
                self._replace_shop_links(db, cursor.lastrowid, payload.get("shop_links") or {})
                self._replace_marketplace_links(db, cursor.lastrowid, payload.get("marketplace_links") or {})
        elif kind == "stock":
            nomenclature = str(payload.get("nomenclature", "")).strip()
            if not nomenclature:
                raise ValueError("Nomenclature is required")
            model = str(payload.get("model", "")).strip() or None
            with self._lock, self.connect() as db:
                cursor = db.execute(
                    "INSERT INTO catalog_items(kind, model, canonical_model, nomenclature, warehouse, quantity, unit_cost_eur, "
                    "origin, paused, created_at, updated_at) VALUES('stock', ?, ?, ?, ?, ?, ?, 'manual', ?, ?, ?)",
                    (
                        model, canonicalize(model) if model else None, nomenclature,
                        str(payload.get("warehouse", "")).strip() or None,
                        number(payload.get("quantity")), number(payload.get("unit_cost_eur")),
                        int(bool(payload.get("paused"))), now, now,
                    ),
                )
        else:
            raise ValueError("kind must be source or stock")
        return self.get_item(cursor.lastrowid)

    def update_item(self, item_id: int, payload: dict[str, Any]) -> dict[str, Any]:
        current = self.get_item(item_id)
        fields: dict[str, Any] = {}
        if current["kind"] == "source":
            if "model" in payload:
                model = str(payload["model"]).strip()
                if not model:
                    raise ValueError("Model is required")
                fields.update(model=model, canonical_model=canonicalize(model))
            if "source_sheets" in payload:
                sheets = [sheet for sheet in payload["source_sheets"] if sheet in SOURCE_SHEETS]
                fields["source_sheets"] = json.dumps(sheets or ["TV"])
        else:
            for name in ("nomenclature", "warehouse"):
                if name in payload:
                    fields[name] = str(payload[name]).strip() or None
            if "model" in payload:
                model = str(payload["model"]).strip() or None
                fields.update(model=model, canonical_model=canonicalize(model) if model else None)
            for name in ("quantity", "unit_cost_eur"):
                if name in payload:
                    fields[name] = number(payload[name])
        if "paused" in payload:
            fields["paused"] = int(bool(payload["paused"]))
        has_shop_links = current["kind"] == "source" and "shop_links" in payload
        has_marketplace_links = current["kind"] == "source" and "marketplace_links" in payload
        if not fields and not has_shop_links and not has_marketplace_links:
            return current
        with self._lock, self.connect() as db:
            if current["kind"] == "source" and "canonical_model" in fields:
                duplicate = db.execute(
                    "SELECT id FROM catalog_items WHERE kind='source' AND canonical_model=? "
                    "AND deleted_at IS NULL AND id<>?",
                    (fields["canonical_model"], item_id),
                ).fetchone()
                if duplicate:
                    raise ValueError(f"Model already exists (item {duplicate['id']})")
            if fields:
                fields["origin"] = "manual"
                fields["updated_at"] = utc_now()
                assignments = ", ".join(f"{name}=?" for name in fields)
                db.execute(f"UPDATE catalog_items SET {assignments} WHERE id=?", (*fields.values(), item_id))
            if has_shop_links:
                self._replace_shop_links(db, item_id, payload.get("shop_links") or {})
            if has_marketplace_links:
                self._replace_marketplace_links(db, item_id, payload.get("marketplace_links") or {})
        return self.get_item(item_id)

    def trash_item(self, item_id: int) -> dict[str, Any]:
        with self._lock, self.connect() as db:
            cursor = db.execute(
                "UPDATE catalog_items SET deleted_at=?, updated_at=? WHERE id=? AND deleted_at IS NULL",
                (utc_now(), utc_now(), item_id),
            )
            if not cursor.rowcount:
                raise KeyError(item_id)
        return self.get_item(item_id)

    def restore_item(self, item_id: int) -> dict[str, Any]:
        with self._lock, self.connect() as db:
            cursor = db.execute(
                "UPDATE catalog_items SET deleted_at=NULL, updated_at=? WHERE id=? AND deleted_at IS NOT NULL",
                (utc_now(), item_id),
            )
            if not cursor.rowcount:
                raise KeyError(item_id)
        return self.get_item(item_id)

    def delete_item_permanently(self, item_id: int) -> None:
        """Permanently remove an item that has already been moved to trash."""
        with self._lock, self.connect() as db:
            item = db.execute(
                "SELECT id FROM catalog_items WHERE id=? AND deleted_at IS NOT NULL", (item_id,)
            ).fetchone()
            if not item:
                raise KeyError(item_id)
            # Historical observations reference catalog items without ON DELETE CASCADE.
            db.execute("DELETE FROM shop_observations WHERE item_id=?", (item_id,))
            db.execute("DELETE FROM assisted_marketplace_observations WHERE item_id=?", (item_id,))
            db.execute("DELETE FROM catalog_items WHERE id=?", (item_id,))

    def promote_stock_item(self, item_id: int) -> dict[str, Any]:
        """Create, restore, or reactivate a monitoring model from a stock item."""
        stock = self.get_item(item_id)
        if stock["kind"] != "stock" or stock["deleted_at"]:
            raise KeyError(item_id)
        model = str(stock.get("model") or "").strip() or infer_model_from_nomenclature(stock.get("nomenclature") or "")
        if not model:
            raise ValueError("No model could be recognized in this stock item; edit the linked model first")
        canonical = canonicalize(model)
        now = utc_now()
        with self._lock, self.connect() as db:
            if not stock.get("model"):
                db.execute(
                    "UPDATE catalog_items SET model=?, canonical_model=?, origin='manual', updated_at=? WHERE id=?",
                    (model, canonical, now, item_id),
                )
            existing = db.execute(
                "SELECT id, deleted_at FROM catalog_items WHERE kind='source' AND canonical_model=? "
                "ORDER BY deleted_at IS NULL DESC, id DESC LIMIT 1",
                (canonical,),
            ).fetchone()
            if existing:
                db.execute(
                    "UPDATE catalog_items SET deleted_at=NULL, paused=0, updated_at=? WHERE id=?",
                    (now, existing["id"]),
                )
                source_id = int(existing["id"])
            else:
                cursor = db.execute(
                    "INSERT INTO catalog_items(kind, model, canonical_model, source_sheets, origin, paused, created_at, updated_at) "
                    "VALUES('source', ?, ?, ?, 'manual', 0, ?, ?)",
                    (model, canonical, json.dumps([source_sheet_for_nomenclature(stock.get("nomenclature") or "")]), now, now),
                )
                source_id = int(cursor.lastrowid)
        return self.get_item(source_id)

    def import_source_workbook(self, path: Path, filename: str) -> dict[str, Any]:
        workbook = load_workbook(path, data_only=True, read_only=True)
        parsed: dict[str, dict[str, Any]] = {}
        for sheet_name in SOURCE_SHEETS:
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, ())
            model_index = next(
                (index for index, value in enumerate(header) if str(value or "").strip().lower() == "model"),
                None,
            )
            if model_index is None:
                continue
            for row in rows:
                if model_index >= len(row) or row[model_index] in (None, ""):
                    continue
                model = str(row[model_index]).strip()
                canonical = canonicalize(model)
                entry = parsed.setdefault(canonical, {"model": model, "source_sheets": []})
                if sheet_name not in entry["source_sheets"]:
                    entry["source_sheets"].append(sheet_name)
        if not parsed:
            raise ValueError("No models found in TV, SB or Monitors sheets")

        inserted = 0
        updated = 0
        now = utc_now()
        with self._lock, self.connect() as db:
            for canonical, item in parsed.items():
                existing = db.execute(
                    "SELECT id FROM catalog_items WHERE kind='source' AND canonical_model=? AND deleted_at IS NULL",
                    (canonical,),
                ).fetchone()
                if existing:
                    db.execute(
                        "UPDATE catalog_items SET model=?, source_sheets=?, source_filename=?, updated_at=? WHERE id=?",
                        (item["model"], json.dumps(item["source_sheets"]), filename, now, existing["id"]),
                    )
                    updated += 1
                else:
                    db.execute(
                        "INSERT INTO catalog_items(kind, model, canonical_model, source_sheets, source_filename, origin, created_at, updated_at) "
                        "VALUES('source', ?, ?, ?, ?, 'file', ?, ?)",
                        (item["model"], canonical, json.dumps(item["source_sheets"]), filename, now, now),
                    )
                    inserted += 1
            self._set_meta(db, "source_upload", {"filename": filename, "uploaded_at": now, "count": len(parsed)})
        return {"imported": len(parsed), "inserted": inserted, "updated": updated}

    def import_stock_workbook(self, path: Path, filename: str) -> dict[str, Any]:
        workbook = load_workbook(path, data_only=True, read_only=True)
        source_models = [item["model"] for item in self.list_items("source", "active")]
        source_models.sort(key=len, reverse=True)
        parsed: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            warehouse: str | None = None
            for row in sheet.iter_rows(values_only=True):
                name = str(row[1]).strip() if len(row) > 1 and row[1] not in (None, "") else ""
                qty = row[2] if len(row) > 2 else None
                cost = row[3] if len(row) > 3 else None
                if not name:
                    continue
                if qty in (None, "") and cost in (None, ""):
                    warehouse = name
                    continue
                if name.lower() in {"item", "total", "nomenclature", "наименование"}:
                    continue
                matched = next((model for model in source_models if canonicalize(model) in canonicalize(name)), None)
                parsed.append(
                    {
                        "nomenclature": name,
                        "model": matched,
                        "canonical_model": canonicalize(matched) if matched else None,
                        "warehouse": warehouse,
                        "quantity": number(qty),
                        "unit_cost_eur": number(cost),
                    }
                )
        if not parsed:
            raise ValueError("No stock items found in columns B-D")

        now = utc_now()
        with self._lock, self.connect() as db:
            db.execute("DELETE FROM catalog_items WHERE kind='stock' AND origin IN ('file', 'v3')")
            for item in parsed:
                db.execute(
                    "INSERT INTO catalog_items(kind, model, canonical_model, nomenclature, warehouse, quantity, unit_cost_eur, "
                    "source_filename, origin, created_at, updated_at) VALUES('stock', ?, ?, ?, ?, ?, ?, ?, 'file', ?, ?)",
                    (
                        item["model"], item["canonical_model"], item["nomenclature"], item["warehouse"],
                        item["quantity"], item["unit_cost_eur"], filename, now, now,
                    ),
                )
            self._set_meta(db, "stock_upload", {"filename": filename, "uploaded_at": now, "count": len(parsed)})
        return {"imported": len(parsed), "matched": sum(1 for item in parsed if item["canonical_model"])}

    def stats(self) -> dict[str, Any]:
        with self.connect() as db:
            rows = db.execute(
                "SELECT kind, SUM(deleted_at IS NULL) active_or_paused, "
                "SUM(deleted_at IS NULL AND paused=0) active, SUM(deleted_at IS NULL AND paused=1) paused, "
                "SUM(deleted_at IS NOT NULL) trash FROM catalog_items GROUP BY kind"
            ).fetchall()
            active_models = {
                row[0] for row in db.execute(
                    "SELECT canonical_model FROM catalog_items WHERE kind='source' AND deleted_at IS NULL AND paused=0"
                )
            }
            stock_models = [
                row[0] for row in db.execute(
                    "SELECT canonical_model FROM catalog_items WHERE kind='stock' AND deleted_at IS NULL AND paused=0"
                )
            ]
        result = {
            "source": {"active_or_paused": 0, "active": 0, "paused": 0, "trash": 0},
            "stock": {"active_or_paused": 0, "active": 0, "paused": 0, "trash": 0},
        }
        for row in rows:
            result[row["kind"]] = {key: int(row[key] or 0) for key in ("active_or_paused", "active", "paused", "trash")}
        result["stock"]["unmatched"] = sum(1 for model in stock_models if not model or model not in active_models)
        return result

    def list_sources(self) -> list[dict[str, Any]]:
        masters = self.get_meta("source_masters", {"marketplace": True, "shop": True})
        with self.connect() as db:
            rows = db.execute("SELECT * FROM monitoring_sources ORDER BY sort_order").fetchall()
        result = []
        for row in rows:
            item = dict(row)
            item["enabled"] = bool(item["enabled"])
            item["master_enabled"] = bool(masters.get(item["kind"], True))
            item["effective_enabled"] = item["enabled"] and item["master_enabled"]
            result.append(item)
        return result

    def update_source(
        self,
        key: str,
        enabled: bool | None = None,
        collection_method: str | None = None,
    ) -> dict[str, Any]:
        with self._lock, self.connect() as db:
            source = db.execute("SELECT kind FROM monitoring_sources WHERE key=?", (key,)).fetchone()
            if not source:
                raise KeyError(key)
            if collection_method is not None:
                allowed = (
                    {"auto", "legacy"}
                    if source["kind"] == "marketplace"
                    else {"auto", "direct", "background", "playwright", "extension", "manual"}
                )
                if collection_method not in allowed:
                    raise ValueError(f"Unsupported {source['kind']} collection method: {collection_method}")
                db.execute(
                    "UPDATE monitoring_sources SET collection_method=? WHERE key=?",
                    (collection_method, key),
                )
            if enabled is not None:
                db.execute("UPDATE monitoring_sources SET enabled=? WHERE key=?", (int(enabled), key))
        return next(item for item in self.list_sources() if item["key"] == key)

    def remember_source_method(self, key: str, collection_method: str) -> None:
        if collection_method not in {"direct", "background", "playwright", "extension"}:
            return
        with self._lock, self.connect() as db:
            db.execute(
                "UPDATE monitoring_sources SET last_success_method=? WHERE key=?",
                (collection_method, key),
            )

    def update_source_master(self, kind: str, enabled: bool) -> dict[str, Any]:
        if kind not in {"marketplace", "shop"}:
            raise ValueError("kind must be marketplace or shop")
        masters = self.get_meta("source_masters", {"marketplace": True, "shop": True})
        masters[kind] = bool(enabled)
        with self._lock, self.connect() as db:
            self._set_meta(db, "source_masters", masters)
        return {"kind": kind, "enabled": bool(enabled)}

    def register_monitoring_session(self, run_id: str, legacy_started: bool, marketplace_keys: list[str]) -> None:
        with self._lock, self.connect() as db:
            db.execute(
                "INSERT OR REPLACE INTO monitoring_sessions(run_id, legacy_started, marketplace_keys, created_at) "
                "VALUES(?, ?, ?, ?)",
                (run_id, int(legacy_started), json.dumps(marketplace_keys), utc_now()),
            )

    def stop_monitoring_session(self, run_id: str) -> None:
        with self._lock, self.connect() as db:
            cursor = db.execute(
                "UPDATE monitoring_sessions SET stopped_at=? WHERE run_id=?",
                (utc_now(), run_id),
            )
            if not cursor.rowcount:
                raise KeyError(run_id)

    def clear_monitoring_session(self, run_id: str) -> None:
        """Clear the current results while retaining a durable empty-table marker."""
        now = utc_now()
        with self._lock, self.connect() as db:
            cursor = db.execute(
                "UPDATE monitoring_sessions SET stopped_at=COALESCE(stopped_at,?),cleared_at=? "
                "WHERE run_id=?",
                (now, now, run_id),
            )
            if not cursor.rowcount:
                raise KeyError(run_id)
            db.execute("DELETE FROM shop_observations WHERE run_id=?", (run_id,))
            db.execute("DELETE FROM assisted_marketplace_observations WHERE run_id=?", (run_id,))

    def monitoring_session(self, run_id: str) -> dict[str, Any] | None:
        with self.connect() as db:
            row = db.execute("SELECT * FROM monitoring_sessions WHERE run_id=?", (run_id,)).fetchone()
        if not row:
            return None
        result = dict(row)
        result["legacy_started"] = bool(result["legacy_started"])
        result["marketplace_keys"] = json.loads(result["marketplace_keys"] or "[]")
        return result

    def latest_monitoring_session(self) -> dict[str, Any] | None:
        with self.connect() as db:
            row = db.execute("SELECT run_id FROM monitoring_sessions ORDER BY created_at DESC LIMIT 1").fetchone()
        return self.monitoring_session(row["run_id"]) if row else None

    def list_monitoring_sessions(self, limit: int = 20) -> list[dict[str, Any]]:
        with self.connect() as db:
            rows = db.execute(
                "SELECT s.*,"
                "(SELECT COUNT(*) FROM shop_observations o WHERE o.run_id=s.run_id) shop_checks,"
                "(SELECT COUNT(*) FROM assisted_marketplace_observations o WHERE o.run_id=s.run_id) assisted_checks,"
                "(SELECT COUNT(*) FROM shop_observations o WHERE o.run_id=s.run_id "
                " AND o.status IN ('PENDING','RUNNING')) shop_pending,"
                "(SELECT COUNT(*) FROM assisted_marketplace_observations o WHERE o.run_id=s.run_id "
                " AND o.status IN ('PENDING','RUNNING')) assisted_pending "
                "FROM monitoring_sessions s ORDER BY s.created_at DESC LIMIT ?",
                (max(1, min(int(limit), 100)),),
            ).fetchall()
        sessions: list[dict[str, Any]] = []
        for row in rows:
            item = dict(row)
            item["legacy_started"] = bool(item["legacy_started"])
            item["marketplace_keys"] = json.loads(item.get("marketplace_keys") or "[]")
            if item.get("cleared_at"):
                item["status"] = "CLEARED"
            elif item.get("stopped_at"):
                item["status"] = "INCOMPLETE"
            elif item.get("shop_pending") or item.get("assisted_pending"):
                item["status"] = "RUNNING"
            else:
                item["status"] = "COMPLETE"
            sessions.append(item)
        return sessions

    def start_shop_run(
        self,
        run_id: str,
        models: list[dict[str, Any]],
        shops: list[dict[str, Any]],
        cache_ttl_seconds: float = 0,
    ) -> None:
        now = datetime.now(UTC)
        cutoff = (now - timedelta(seconds=max(0, cache_ttl_seconds))).isoformat()
        now_text = now.isoformat()
        with self._lock, self.connect() as db:
            db.execute("DELETE FROM shop_observations WHERE run_id=?", (run_id,))
            for model in models:
                for shop in shops:
                    protection = db.execute(
                        "SELECT cooldown_until, reason FROM shop_protection_state "
                        "WHERE shop_key=? AND cooldown_until>?",
                        (shop["key"], now_text),
                    ).fetchone()
                    if protection:
                        db.execute(
                            "INSERT INTO shop_observations(run_id,item_id,shop_key,status,error,retry_after) "
                            "VALUES(?,?,?,'COOLDOWN',?,?)",
                            (run_id, model["id"], shop["key"], protection["reason"], protection["cooldown_until"]),
                        )
                        continue
                    cached = None
                    if cache_ttl_seconds > 0:
                        cached = db.execute(
                            "SELECT title,price_eur,availability,product_url,search_url,checked_at,"
                            "collection_method,attempts_json "
                            "FROM shop_observations WHERE run_id<>? AND item_id=? AND shop_key=? "
                            "AND status='SUCCESS' AND price_eur IS NOT NULL AND checked_at>=? "
                            "ORDER BY checked_at DESC LIMIT 1",
                            (run_id, model["id"], shop["key"], cutoff),
                        ).fetchone()
                    if cached:
                        db.execute(
                            "INSERT INTO shop_observations(run_id,item_id,shop_key,status,title,price_eur,availability,"
                            "product_url,search_url,checked_at,collection_method,attempts_json,cached) "
                            "VALUES(?,?,?,'SUCCESS',?,?,?,?,?,?,?,?,1)",
                            (
                                run_id, model["id"], shop["key"], cached["title"], cached["price_eur"],
                                cached["availability"], cached["product_url"], cached["search_url"], cached["checked_at"],
                                cached["collection_method"], cached["attempts_json"] or "[]",
                            ),
                        )
                    else:
                        db.execute(
                            "INSERT INTO shop_observations(run_id, item_id, shop_key, status) VALUES(?, ?, ?, 'PENDING')",
                            (run_id, model["id"], shop["key"]),
                        )

    def finish_shop_observation(
        self,
        run_id: str,
        item_id: int,
        shop_key: str,
        status: str,
        *,
        title: str | None = None,
        price_eur: float | None = None,
        availability: str | None = None,
        product_url: str | None = None,
        search_url: str | None = None,
        error: str | None = None,
        retry_after: str | None = None,
        collection_method: str | None = None,
        attempts: list[dict[str, Any]] | None = None,
    ) -> None:
        with self._lock, self.connect() as db:
            db.execute(
                "UPDATE shop_observations SET status=?, title=?, price_eur=?, availability=?, product_url=?, "
                "search_url=?, error=?, checked_at=?, retry_after=?, collection_method=?, attempts_json=?, cached=0 "
                "WHERE run_id=? AND item_id=? AND shop_key=?",
                (
                    status, title, price_eur, availability, product_url, search_url, error, utc_now(), retry_after,
                    collection_method, json.dumps(attempts or [], ensure_ascii=False),
                    run_id, item_id, shop_key,
                ),
            )

    def retry_shop_observation(self, run_id: str, item_id: int, shop_key: str) -> None:
        with self._lock, self.connect() as db:
            cursor = db.execute(
                "UPDATE shop_observations SET status='PENDING', title=NULL, price_eur=NULL, availability=NULL, "
                "product_url=NULL, error=NULL, checked_at=NULL, retry_after=NULL, collection_method=NULL, "
                "attempts_json='[]', cached=0 "
                "WHERE run_id=? AND item_id=? AND shop_key=?",
                (run_id, item_id, shop_key),
            )
            if not cursor.rowcount:
                raise KeyError((run_id, item_id, shop_key))
            db.execute("DELETE FROM shop_protection_state WHERE shop_key=?", (shop_key,))

    def ensure_shop_observation(self, run_id: str, item_id: int, shop_key: str) -> None:
        with self._lock, self.connect() as db:
            db.execute(
                "INSERT OR IGNORE INTO shop_observations(run_id,item_id,shop_key,status) "
                "VALUES(?,?,?,'PENDING')", (run_id, item_id, shop_key),
            )

    def resolve_shop_observation(
        self,
        run_id: str,
        item_id: int,
        shop_key: str,
        status: str,
        *,
        price_eur: float | None = None,
        availability: str | None = None,
    ) -> dict[str, Any]:
        if status not in {"NOT_FOUND", "SUCCESS"}:
            raise ValueError("Manual shop status must be NOT_FOUND or SUCCESS")
        if status == "SUCCESS" and (price_eur is None or price_eur < 0):
            raise ValueError("A non-negative price is required for manual success")
        attempt = [{"method": "manual", "result": "CONFIRMED", "duration_ms": 0}]
        with self._lock, self.connect() as db:
            current = db.execute(
                "SELECT product_url,search_url FROM shop_observations "
                "WHERE run_id=? AND item_id=? AND shop_key=?",
                (run_id, item_id, shop_key),
            ).fetchone()
            if not current:
                raise KeyError((run_id, item_id, shop_key))
            db.execute(
                "UPDATE shop_observations SET status=?, price_eur=?, availability=?, error=NULL, checked_at=?, "
                "retry_after=NULL, cached=0, collection_method='manual', attempts_json=? "
                "WHERE run_id=? AND item_id=? AND shop_key=?",
                (
                    status,
                    price_eur if status == "SUCCESS" else None,
                    (availability or "IN_STOCK") if status == "SUCCESS" else None,
                    utc_now(),
                    json.dumps(attempt),
                    run_id,
                    item_id,
                    shop_key,
                ),
            )
            row = db.execute(
                "SELECT * FROM shop_observations WHERE run_id=? AND item_id=? AND shop_key=?",
                (run_id, item_id, shop_key),
            ).fetchone()
        result = dict(row)
        result["cached"] = bool(result.get("cached"))
        result["attempts"] = json.loads(result.pop("attempts_json", "[]") or "[]")
        return result

    def shop_observation_pending(self, run_id: str, item_id: int, shop_key: str) -> bool:
        with self.connect() as db:
            row = db.execute(
                "SELECT status FROM shop_observations WHERE run_id=? AND item_id=? AND shop_key=?",
                (run_id, item_id, shop_key),
            ).fetchone()
        return bool(row and row["status"] == "PENDING")

    def pause_shop_for_protection(
        self, run_id: str, shop_key: str, cooldown_seconds: float, reason: str
    ) -> str:
        now = datetime.now(UTC)
        cooldown_until = (now + timedelta(seconds=max(1, cooldown_seconds))).isoformat()
        message = str(reason or "Shop protection requested a cooldown")[:500]
        with self._lock, self.connect() as db:
            db.execute(
                "INSERT INTO shop_protection_state(shop_key,cooldown_until,reason,updated_at) VALUES(?,?,?,?) "
                "ON CONFLICT(shop_key) DO UPDATE SET cooldown_until=excluded.cooldown_until, "
                "reason=excluded.reason, updated_at=excluded.updated_at",
                (shop_key, cooldown_until, message, now.isoformat()),
            )
            db.execute(
                "UPDATE shop_observations SET status='COOLDOWN', error=?, retry_after=? "
                "WHERE run_id=? AND shop_key=? AND status='PENDING'",
                (message, cooldown_until, run_id, shop_key),
            )
        return cooldown_until

    def cancel_shop_run(self, run_id: str) -> int:
        """Mark unfinished direct-shop observations as stopped by the user."""
        with self._lock, self.connect() as db:
            cursor = db.execute(
                "UPDATE shop_observations SET status='INCOMPLETE', error='Stopped by user', checked_at=? "
                "WHERE run_id=? AND status IN ('PENDING', 'RUNNING')",
                (utc_now(), run_id),
            )
            return int(cursor.rowcount)

    def cancel_legacy_run(self, legacy_database: Path, run_id: str) -> int:
        """Finalize a run left active inside the v3 marketplace engine."""
        if not legacy_database.exists():
            return 0
        now = utc_now()
        with self._lock, sqlite3.connect(legacy_database, timeout=30) as db:
            run_cursor = db.execute(
                "UPDATE monitoring_runs SET status='INCOMPLETE', finished_at=?, error='Stopped by user' "
                "WHERE id=? AND status='RUNNING'",
                (now, run_id),
            )
            db.execute(
                "UPDATE marketplace_tasks SET status='INCOMPLETE', finished_at=?, error='Stopped by user' "
                "WHERE run_id=? AND status IN ('PENDING', 'RUNNING')",
                (now, run_id),
            )
            return int(run_cursor.rowcount)

    def shop_run(self, run_id: str) -> dict[str, Any]:
        with self.connect() as db:
            rows = db.execute(
                "SELECT o.*, i.model, s.name shop_name, s.country FROM shop_observations o "
                "JOIN catalog_items i ON i.id=o.item_id JOIN monitoring_sources s ON s.key=o.shop_key "
                "WHERE o.run_id=? ORDER BY i.model, s.sort_order",
                (run_id,),
            ).fetchall()
        results = [dict(row) for row in rows]
        for result in results:
            result["cached"] = bool(result.get("cached"))
            result["attempts"] = json.loads(result.pop("attempts_json", "[]") or "[]")
        pending = sum(1 for row in results if row["status"] == "PENDING")
        return {"status": "RUNNING" if pending else "COMPLETE", "pending": pending, "results": results}

    def start_assisted_marketplace_run(
        self,
        run_id: str,
        models: list[dict[str, Any]],
        marketplaces: list[dict[str, Any]],
        cache_ttl_seconds: float = 0,
    ) -> None:
        cutoff = (datetime.now(UTC) - timedelta(seconds=max(0, cache_ttl_seconds))).isoformat()
        with self._lock, self.connect() as db:
            db.execute("DELETE FROM assisted_marketplace_observations WHERE run_id=?", (run_id,))
            for model in models:
                for marketplace in marketplaces:
                    cached = None
                    if cache_ttl_seconds > 0:
                        cached = db.execute(
                            "SELECT title,seller_name,price_eur,availability,product_url,search_url,checked_at,"
                            "collection_method,attempts_json FROM assisted_marketplace_observations "
                            "WHERE run_id<>? AND item_id=? AND marketplace_key=? AND status='SUCCESS' "
                            "AND price_eur IS NOT NULL AND checked_at>=? ORDER BY checked_at DESC LIMIT 1",
                            (run_id, model["id"], marketplace["key"], cutoff),
                        ).fetchone()
                    if cached:
                        db.execute(
                            "INSERT INTO assisted_marketplace_observations("
                            "run_id,item_id,marketplace_key,status,title,seller_name,price_eur,availability,"
                            "product_url,search_url,checked_at,cached,collection_method,attempts_json) "
                            "VALUES(?,?,?,'SUCCESS',?,?,?,?,?,?,?,1,?,?)",
                            (
                                run_id, model["id"], marketplace["key"], cached["title"],
                                cached["seller_name"], cached["price_eur"], cached["availability"],
                                cached["product_url"], cached["search_url"], cached["checked_at"],
                                cached["collection_method"], cached["attempts_json"] or "[]",
                            ),
                        )
                    else:
                        db.execute(
                            "INSERT INTO assisted_marketplace_observations("
                            "run_id,item_id,marketplace_key,status) VALUES(?,?,?,'PENDING')",
                            (run_id, model["id"], marketplace["key"]),
                        )

    def finish_assisted_marketplace_observation(
        self,
        run_id: str,
        item_id: int,
        marketplace_key: str,
        status: str,
        *,
        title: str | None = None,
        seller_name: str | None = None,
        price_eur: float | None = None,
        availability: str | None = None,
        product_url: str | None = None,
        search_url: str | None = None,
        error: str | None = None,
        collection_method: str | None = "extension",
        attempts: list[dict[str, Any]] | None = None,
    ) -> None:
        with self._lock, self.connect() as db:
            db.execute(
                "UPDATE assisted_marketplace_observations SET status=?,title=?,seller_name=?,price_eur=?,"
                "availability=?,product_url=?,search_url=?,error=?,checked_at=?,cached=0,collection_method=?,"
                "attempts_json=? WHERE run_id=? AND item_id=? AND marketplace_key=?",
                (
                    status, title, seller_name, price_eur, availability, product_url, search_url, error,
                    utc_now(), collection_method, json.dumps(attempts or [], ensure_ascii=False),
                    run_id, item_id, marketplace_key,
                ),
            )

    def retry_assisted_marketplace_observation(
        self, run_id: str, item_id: int, marketplace_key: str
    ) -> None:
        with self._lock, self.connect() as db:
            cursor = db.execute(
                "UPDATE assisted_marketplace_observations SET status='PENDING',title=NULL,seller_name=NULL,"
                "price_eur=NULL,availability=NULL,product_url=NULL,error=NULL,checked_at=NULL,cached=0,"
                "collection_method=NULL,attempts_json='[]' WHERE run_id=? AND item_id=? AND marketplace_key=?",
                (run_id, item_id, marketplace_key),
            )
            if not cursor.rowcount:
                raise KeyError((run_id, item_id, marketplace_key))

    def ensure_assisted_marketplace_observation(
        self, run_id: str, item_id: int, marketplace_key: str
    ) -> None:
        with self._lock, self.connect() as db:
            db.execute(
                "INSERT OR IGNORE INTO assisted_marketplace_observations("
                "run_id,item_id,marketplace_key,status) VALUES(?,?,?,'PENDING')",
                (run_id, item_id, marketplace_key),
            )

    def resolve_assisted_marketplace_observation(
        self,
        run_id: str,
        item_id: int,
        marketplace_key: str,
        status: str,
        *,
        price_eur: float | None = None,
        availability: str | None = None,
        seller_name: str | None = None,
    ) -> dict[str, Any]:
        if status not in {"NOT_FOUND", "SUCCESS"}:
            raise ValueError("Manual marketplace status must be NOT_FOUND or SUCCESS")
        if status == "SUCCESS" and (price_eur is None or price_eur < 0):
            raise ValueError("A non-negative price is required for manual success")
        attempts = [{"method": "manual", "result": "CONFIRMED", "duration_ms": 0}]
        with self._lock, self.connect() as db:
            cursor = db.execute(
                "UPDATE assisted_marketplace_observations SET status=?,seller_name=?,"
                "price_eur=?,availability=?,error=NULL,checked_at=?,cached=0,collection_method='manual',"
                "attempts_json=? WHERE run_id=? AND item_id=? AND marketplace_key=?",
                (
                    status, (str(seller_name or "").strip() or "Manual verification") if status == "SUCCESS" else None,
                    price_eur if status == "SUCCESS" else None,
                    (availability or "IN_STOCK") if status == "SUCCESS" else None,
                    utc_now(), json.dumps(attempts), run_id, item_id, marketplace_key,
                ),
            )
            if not cursor.rowcount:
                raise KeyError((run_id, item_id, marketplace_key))
            row = db.execute(
                "SELECT * FROM assisted_marketplace_observations WHERE run_id=? AND item_id=? "
                "AND marketplace_key=?", (run_id, item_id, marketplace_key),
            ).fetchone()
        return dict(row)

    def assisted_marketplace_observation_pending(
        self, run_id: str, item_id: int, marketplace_key: str
    ) -> bool:
        with self.connect() as db:
            row = db.execute(
                "SELECT status FROM assisted_marketplace_observations WHERE run_id=? AND item_id=? "
                "AND marketplace_key=?", (run_id, item_id, marketplace_key),
            ).fetchone()
        return bool(row and row["status"] == "PENDING")

    def cancel_assisted_marketplace_run(self, run_id: str) -> int:
        with self._lock, self.connect() as db:
            cursor = db.execute(
                "UPDATE assisted_marketplace_observations SET status='INCOMPLETE',"
                "error='Stopped by user',checked_at=? WHERE run_id=? AND status IN ('PENDING','RUNNING')",
                (utc_now(), run_id),
            )
            return int(cursor.rowcount)

    def assisted_marketplace_run(self, run_id: str) -> dict[str, Any]:
        with self.connect() as db:
            rows = db.execute(
                "SELECT o.*,i.model,i.canonical_model,s.name marketplace_name,s.country "
                "FROM assisted_marketplace_observations o JOIN catalog_items i ON i.id=o.item_id "
                "JOIN monitoring_sources s ON s.key=o.marketplace_key WHERE o.run_id=? "
                "ORDER BY i.model,s.sort_order", (run_id,),
            ).fetchall()
        results: list[dict[str, Any]] = []
        for row in rows:
            item = dict(row)
            attempt_details = json.loads(item.pop("attempts_json", "[]") or "[]")
            offer = None
            if item["status"] == "SUCCESS" and item.get("price_eur") is not None:
                offer = {
                    "price_eur": item["price_eur"],
                    "store": item.get("seller_name") or item["marketplace_name"],
                    "url": item.get("product_url") or item.get("search_url"),
                }
            results.append({
                **item,
                "id": f"assisted:{run_id}:{item['item_id']}:{item['marketplace_key']}",
                "source_model": item["model"],
                "marketplace": item["marketplace_name"],
                "matched_title": item.get("title"),
                "cheapest_in_stock": offer if item.get("availability") != "PRE_ORDER" else None,
                "cheapest_pre_order": offer if item.get("availability") == "PRE_ORDER" else None,
                "attempts": len(attempt_details),
                "attempt_details": attempt_details,
                "finished_at": item.get("checked_at"),
                "cached": bool(item.get("cached")),
                "assisted": True,
            })
        pending = sum(1 for item in results if item["status"] == "PENDING")
        return {"status": "RUNNING" if pending else "COMPLETE", "pending": pending, "results": results}

    def prepare_legacy(self, original_database: Path, legacy_database: Path, workbook_path: Path) -> None:
        legacy_database.parent.mkdir(parents=True, exist_ok=True)
        if not legacy_database.exists() and original_database.exists():
            source = sqlite3.connect(original_database)
            destination = sqlite3.connect(legacy_database)
            try:
                source.backup(destination)
            finally:
                source.close()
                destination.close()
        self._ensure_legacy_schema(legacy_database)
        self._write_active_workbook(workbook_path)
        self._write_legacy_stock(legacy_database)

    def _ensure_legacy_schema(self, database: Path) -> None:
        with sqlite3.connect(database) as db:
            db.executescript(
                """
                CREATE TABLE IF NOT EXISTS monitoring_runs (
                    id TEXT PRIMARY KEY, trigger TEXT NOT NULL, status TEXT NOT NULL,
                    started_at TEXT NOT NULL, finished_at TEXT, source_workbook TEXT,
                    export_path TEXT, error TEXT
                );
                CREATE TABLE IF NOT EXISTS products (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, canonical_model TEXT UNIQUE NOT NULL,
                    model TEXT NOT NULL, source_sheets TEXT NOT NULL, updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS marketplace_tasks (
                    id TEXT PRIMARY KEY, run_id TEXT NOT NULL REFERENCES monitoring_runs(id),
                    canonical_model TEXT NOT NULL, marketplace TEXT NOT NULL, status TEXT NOT NULL,
                    attempts INTEGER NOT NULL DEFAULT 0, product_url TEXT, error TEXT,
                    started_at TEXT, finished_at TEXT, matched_title TEXT
                );
                CREATE TABLE IF NOT EXISTS offer_observations (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, task_id TEXT NOT NULL REFERENCES marketplace_tasks(id),
                    store_name TEXT NOT NULL, original_price TEXT NOT NULL, currency TEXT NOT NULL,
                    price_eur TEXT NOT NULL, availability TEXT NOT NULL, product_url TEXT NOT NULL,
                    checked_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS application_logs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, run_id TEXT, task_id TEXT, level TEXT NOT NULL,
                    message TEXT NOT NULL, context TEXT NOT NULL, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS stock_items (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, canonical_model TEXT, nomenclature TEXT NOT NULL,
                    warehouse TEXT, quantity TEXT NOT NULL, unit_cost_eur TEXT NOT NULL,
                    source_filename TEXT NOT NULL, uploaded_at TEXT NOT NULL
                );
                """
            )

    def _write_active_workbook(self, path: Path) -> None:
        items = self.list_items("source", "active")
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheets = {name: workbook.create_sheet(name) for name in SOURCE_SHEETS}
        for sheet in sheets.values():
            sheet.append(["Model"])
        for item in items:
            if item["paused"] or item["deleted_at"]:
                continue
            targets = item["source_sheets"] or ["TV"]
            for name in targets:
                if name in sheets:
                    sheets[name].append([item["model"]])
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)

    def _write_legacy_stock(self, database: Path) -> None:
        items = self.list_items("stock", "active")
        now = utc_now()
        with sqlite3.connect(database) as db:
            db.execute("DELETE FROM stock_items")
            for item in items:
                if item["paused"] or item["deleted_at"]:
                    continue
                db.execute(
                    "INSERT INTO stock_items(canonical_model, nomenclature, warehouse, quantity, unit_cost_eur, source_filename, uploaded_at) "
                    "VALUES(?, ?, ?, ?, ?, ?, ?)",
                    (
                        item["canonical_model"], item["nomenclature"], item["warehouse"],
                        str(item["quantity"] or 0), str(item["unit_cost_eur"] or 0),
                        item["source_filename"] or "manual-v4", now,
                    ),
                )
